"""Reproducible PPTXer stress decks and a browsable PowerPoint render gallery.

Run with the project's Python. Artifacts stay under artifacts/pptxer-visibility.
Every stress payload is exactly 256 characters, including a unique tail marker.
"""
from __future__ import annotations

import argparse
import json
import os
from pathlib import Path
import sys
from pptxer_visibility_gallery import write_gallery

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "Tlamatini/agent/agents/pptxer"))
import pptxer_audit as audit
import pptxer_build as build
import pptxer_docmodel as model
import pptxer_nuance as nuance
import pptxer_render as render
import pptxer_theme as themes

PAYLOADS = {}


def payload(key, kind="prose"):
    samples = {
        "prose": "Wide windows need measured wrapping. Preserve every word, keep readable type, and continue on another slide when the available space is full. ",
        "wide": "WWWW MMMM WIDE WINDOWS MEASURE EVERY CHARACTER ",
        "token": "W",
        "spanish": "Información técnica: acción, medición y verificación. El contenido completo debe permanecer visible dentro de sus límites. ",
        "cjk": "測定結果を確認して文字が枠内に収まるように配置する。",
    }
    tail = f" END_{key}"
    sample = samples[kind]
    text = (sample * 256)[:256 - len(tail)] + tail
    PAYLOADS[key] = text
    return text


def stress_deck(extended=False):
    PAYLOADS.clear()
    p = payload
    slides = [
        model.Slide(kind="title_slide", title=p("COVER"), subtitle=p("SUBTITLE"), kicker=p("KICKER")),
        model.Slide(kind="bullets", title="Six long bullets", bullets=[p(f"BULLET{i}") for i in range(6)]),
        model.Slide(kind="bullets", title=p("HEADING"), subtitle=p("SUBHEADING"), bullets=[p("BODY")]),
        model.Slide(kind="stats", title="Six cards with 256-character labels", stats=[{"value": f"{i + 1}42%", "label": p(f"STAT{i}")} for i in range(6)]),
        model.Slide(kind="two_column", title="Long column titles and paragraphs", columns=[{"title": p(f"COLTITLE{i}"), "bullets": [p(f"COL{i}_{j}") for j in range(3)]} for i in range(2)]),
        model.Slide(kind="table", title="Native table with long cells", table={"headers": ["Item", "Description", "Evidence"], "rows": [[str(i), p(f"CELL{i}"), p(f"EVIDENCE{i}", "wide")] for i in range(4)]}),
        model.Slide(kind="timeline", title="Six milestones with long descriptions", timeline=[{"when": f"Phase {i + 1}", "what": p(f"EVENT{i}")} for i in range(6)]),
        model.Slide(kind="diagram", title="Flow nodes with long labels", diagram={"kind": "flow", "nodes": [p(f"NODE{i}") for i in range(4)]}),
        model.Slide(kind="quote", quote=p("QUOTE"), attribution=p("ATTRIBUTION")),
        model.Slide(kind="code", title="Code indentation and a long string", code="def verify():\n    text = '" + p("CODE", "token") + "'\n    return text\n"),
        model.Slide(kind="bullets", title="Wide uppercase text", bullets=[p(f"WIDE{i}", "wide") for i in range(4)]),
        model.Slide(kind="bullets", title="Unbroken 256-character identifier", body="W" * 256),
        model.Slide(kind="bullets", title="Accented Spanish", bullets=[p(f"ES{i}", "spanish") for i in range(4)]),
        model.Slide(kind="section_break", title=p("SECTION")),
        model.Slide(kind="statement", body=p("STATEMENT")),
        model.Slide(kind="closing", title=p("CLOSING"), subtitle=p("CLOSINGSUB")),
        model.Slide(kind="image_right", title="Long image-side copy", bullets=[p(f"SIDE{i}") for i in range(6)]),
        model.Slide(kind="stats", title="All eight items must survive", stats=[{"value": str(i), "label": p(f"EXTRA{i}")} for i in range(8)]),
    ]
    if extended:
        for kind in ("funnel", "pyramid", "cycle", "matrix", "hub", "stack", "comparison", "venn"):
            slides.append(model.Slide(kind="diagram", title=f"Long {kind} labels", diagram={"kind": kind,
                         "nodes": [p(f"{kind}{i}") for i in range(3 if kind == "venn" else 4)]}))
        for kind in ("bar", "pie", "doughnut", "line"):
            slides.append(model.Slide(kind="chart", title=f"Long {kind} chart labels", chart={"kind": kind,
                         "categories": [p(f"{kind}CAT{i}") for i in range(3)],
                         "series": [{"name": p(f"{kind}SERIES"), "values": [12, 25, 18]}]}))
        slides.extend([
            model.Slide(kind="stats", title="Long values and labels", stats=[{"value": p("LONGVALUE"), "label": p("LONGLABEL")}]),
            model.Slide(kind="timeline", title="Long milestone names", timeline=[{"when": p(f"WHEN{i}"), "what": p(f"WHAT{i}")} for i in range(3)]),
            model.Slide(kind="agenda", title="Long agenda", bullets=[p(f"AGENDA{i}") for i in range(6)]),
            model.Slide(kind="comparison", title="Long comparison", columns=[{"title": p(f"COMPARETITLE{i}"), "bullets": [p(f"COMPARE{i}")]} for i in range(2)]),
            model.Slide(kind="image_left", title="Image and text", images=["test-pattern"], body=p("IMAGELEFT")),
            model.Slide(kind="image_full", title=p("IMAGEFULLTITLE"), subtitle=p("IMAGEFULLSUB"), images=["test-pattern"]),
            model.Slide(kind="gallery", title=p("GALLERYTITLE"), images=["test-pattern"] * 8),
            model.Slide(kind="media", title="Unavailable media fallback", video="absent.mp4", body=p("MEDIAFALLBACK")),
            model.Slide(kind="bullets", title="Text without word spaces", body=p("CJK", "cjk")),
        ])
    return model.Deck(slides, title="PPTXer visibility stress test")


def check_content(path):
    """Prove full strings survived, including the tail of every long field."""
    from pptx import Presentation
    prs = Presentation(path)
    frames = [frame.text for slide in prs.slides for _, frame, _, _ in audit._text_frames(slide.shapes)]
    normalized = "".join("".join(frames).split()).casefold()
    missing = [key for key, value in PAYLOADS.items()
               if "".join(value.split()).casefold() not in normalized]
    if "w" * 256 not in normalized:
        missing.append("unbroken-identifier")
    return {"payloads": len(PAYLOADS) + 1, "length": 256, "missing": missing}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--label", default="after")
    parser.add_argument("--sizes", nargs="+", default=["16:9", "4:3", "vertical"])
    parser.add_argument("--themes", nargs="+", default=["corporate", "cyberpunk"])
    parser.add_argument("--render", action="store_true")
    parser.add_argument("--existing", action="store_true", help="Re-audit saved decks without rebuilding them")
    parser.add_argument("--allow-defects", action="store_true", help="Record a failing baseline without exiting nonzero")
    parser.add_argument("--extended", action="store_true", help="Also exercise charts, all diagram layouts, and real image placement")
    args = parser.parse_args()
    dest = ROOT / "artifacts/pptxer-visibility" / args.label
    dest.mkdir(parents=True, exist_ok=True)
    os.environ["TLAMATINI_TEMP"] = str(ROOT / "Temp")
    reports = []
    for name in args.themes:
        for size in args.sizes:
            key = f"{name}-{size.replace(':', '-') }"
            folder = dest / key
            folder.mkdir(exist_ok=True)
            deck = stress_deck(args.extended)
            media = {}
            if args.extended:
                from PIL import Image, ImageDraw
                from types import SimpleNamespace
                pattern = Image.new("RGB", (960, 360), "#dbeafe")
                draw = ImageDraw.Draw(pattern)
                for x in range(0, 960, 80):
                    draw.rectangle((x, 0, x + 40, 360), fill="#1d4ed8")
                path = folder / "test-pattern.png"
                pattern.save(path)
                media["test-pattern"] = SimpleNamespace(path=str(path), width=960, height=360, ok=True)
            theme = themes.build_theme(nuance.classify(deck.plain_text(), requested=name), {})
            if args.existing:
                result = build.BuildResult()
                result.path = str(folder / "stress.pptx")
                result.ok = Path(result.path).is_file()
                from pptx import Presentation
                result.slide_count = len(Presentation(result.path).slides)
            else:
                result = build.build_deck(deck, theme, str(folder / "stress.pptx"), {"slide_size": size, "generate_art": False}, media_index=media)
            rr = render.render_slides(result.path, str(folder / "slides"), prefer="powerpoint", timeout=180) if args.render and result.ok else None
            report = audit.audit_pptx(result.path, rr, text_boxes=result.text_boxes).as_dict() if result.ok else {}
            coverage = check_content(result.path) if result.ok else {"missing": ["build failed"]}
            entry = {"case": key, "build": result.as_dict(), "audit": report, "render": rr, "content": coverage}
            reports.append(entry)
            (folder / "report.json").write_text(json.dumps(entry, indent=2, ensure_ascii=False), encoding="utf-8")
            print(json.dumps({"case": key, "ok": result.ok, "slides": result.slide_count, "warnings": len(result.warnings), "diagnostics": len(result.placement_diagnostics), "defects": report.get("defects"), "missing": coverage["missing"], "renderer": (rr or {}).get("tier")}), flush=True)
    (dest / "report.json").write_text(json.dumps(reports, indent=2, ensure_ascii=False), encoding="utf-8")
    write_gallery(dest, reports)
    failed = any(not e["build"]["ok"] or e["build"]["placement_diagnostics"]
                 or not e["audit"].get("layout_clean") or e["content"]["missing"]
                 or (args.render and not e["audit"].get("ground_truth")) for e in reports)
    if failed and not args.allow_defects:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
