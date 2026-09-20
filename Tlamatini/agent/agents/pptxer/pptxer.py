# ══════════════════════════════════════════════════════════════════════════════
#   ✦  T L A M A T I N I  ✦   —   "one who knows"
#
#   Created by  Angela López Mendoza   ·   @angelahack1
#   Developer · Architect · Creator of Tlamatini
#
#   Every line of this file was written by Angela López Mendoza.
# ══════════════════════════════════════════════════════════════════════════════
#   Tlamatini Author Banner — do not remove (releases scrub the name automatically)
#
# PPTXer Agent — Tlamatini's PRESENTATION COMPOSER.
#
# Action: Triggered by upstream -> resolve the content -> decide what KIND of deck
#         it is -> design a complete visual system for it -> place every shape
#         with measured geometry -> embed local AND internet media -> paint its
#         own artwork -> write ONE .pptx -> RE-OPEN IT AND LOOK AT IT -> emit
#         INI_SECTION_PPTXER -> ALWAYS trigger downstream (success OR failure OR
#         fail-safe refusal).
#
# PPTXer AUTHORS presentations; File-Extractor / File-Interpreter READ documents.
# It is the PowerPoint sibling of PDFer (which composes PDFs) and LaTeXer (which
# typesets .tex). ZERO new dependencies: python-pptx + pillow + pymupdf + opencv
# already ship with Tlamatini.
#
# ⚠️ SELF-CONTAINED BY DESIGN (Angela's instruction). Where PPTXer needs a
# capability another agent has — Shoter's screen capture, Googler's fetching —
# that code is written HERE, in this agent's own modules, not borrowed. It never
# calls another pool agent and never imports agent.*, because a pool subprocess
# has no sys.path back into the Django app.
#
# The heavy work lives in flat SIBLING modules next to this file (NOT a package:
# the agent is copied into a runtime directory and run as `python pptxer.py`, so
# sys.path[0] is that directory and a flat neighbour imports reliably in source,
# frozen and self-modify builds — the FlowCreator result_to_flw.py precedent).
# The group import is FAIL-OPEN: a partial copy degrades to a clear refusal
# instead of a crash.

import os
import sys

# FIX: Disable Intel Fortran runtime Ctrl+C handler
os.environ['FOR_DISABLE_CONSOLE_CTRL_HANDLER'] = '1'

# ── Tlamatini Temp policy: temporary files ONLY under <app>/Temp ─────────
# Honor TLAMATINI_TEMP (exported by the Tlamatini core, inherited by every spawned
# agent via get_agent_env's os.environ.copy()) so every intermediate file this
# agent writes — fetched media, generated artwork, rendered slide images — lands
# under <app>/Temp, never C:\Temp / %TEMP% / the OS default. Fail-open when unset.
if (os.environ.get('TLAMATINI_TEMP') or '').strip():
    try:
        import tempfile as _tlt_tempfile
        _tlt_temp_root = os.environ['TLAMATINI_TEMP'].strip()
        os.makedirs(_tlt_temp_root, exist_ok=True)
        _tlt_tempfile.tempdir = _tlt_temp_root
        os.environ['TEMP'] = _tlt_temp_root
        os.environ['TMP'] = _tlt_temp_root
    except Exception:
        pass

import json
import logging
import subprocess
import time

import yaml

# -- conhost.exe orphan guard ------------------------------------------
if os.name == 'nt' and not getattr(subprocess, '_conhost_guard_applied', False):
    _CHG_NO_WINDOW = subprocess.CREATE_NO_WINDOW
    _CHG_RESPECT = (
        _CHG_NO_WINDOW
        | getattr(subprocess, 'CREATE_NEW_CONSOLE', 0)
        | getattr(subprocess, 'DETACHED_PROCESS', 0)
    )
    _chg_orig_init = subprocess.Popen.__init__

    def _chg_guarded_init(self, *args, **kwargs):
        cf = kwargs.get('creationflags', 0) or 0
        if not (cf & _CHG_RESPECT):
            kwargs['creationflags'] = cf | _CHG_NO_WINDOW
        return _chg_orig_init(self, *args, **kwargs)

    subprocess.Popen.__init__ = _chg_guarded_init
    subprocess._conhost_guard_applied = True

# Set working directory to script location
try:
    script_dir = os.path.dirname(os.path.abspath(__file__))
    os.chdir(script_dir)
except Exception as e:
    sys.stderr.write(f"Critical Error: Failed to set working directory: {e}\n")

CURRENT_DIR_NAME = os.path.basename(os.path.dirname(os.path.abspath(__file__)))
LOG_FILE_PATH = f"{CURRENT_DIR_NAME}.log"

# Reanimation detection: AGENT_REANIMATED=1 means resume from pause
_IS_REANIMATED = os.environ.get('AGENT_REANIMATED') == '1'
if not _IS_REANIMATED:
    open(LOG_FILE_PATH, 'w').close()

logging.basicConfig(
    filename=LOG_FILE_PATH,
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    encoding='utf-8'
)
console_handler = logging.StreamHandler()
console_handler.setLevel(logging.INFO)
console_handler.setFormatter(logging.Formatter('%(asctime)s - %(levelname)s - %(message)s'))
logging.getLogger().addHandler(console_handler)

# ========================================
# HELPER FUNCTIONS (copied verbatim from the shared pool-agent boilerplate)
# ========================================


def _load_config_file(path: str = "config.yaml") -> dict:
    try:
        with open(path, 'r', encoding='utf-8') as f:
            return yaml.safe_load(f) or {}
    except FileNotFoundError:
        logging.error(f"❌ Config file not found: {path}")
        return {}
    except Exception as e:
        logging.error(f"❌ Failed to read config: {e}")
        return {}


def load_config(*args, **kwargs):
    """Apply Config -> Models choices while retaining explicit agent overrides."""
    import importlib.util as _model_import
    from pathlib import Path as _ModelPath
    config = _load_config_file(*args, **kwargs)
    here = _ModelPath(__file__).resolve()
    candidates = [here.parent / 'model_settings.py']
    candidates += [p / 'model_settings.py' for p in here.parents if p.name == 'agents']
    for shared in candidates:
        if shared.is_file():
            spec = _model_import.spec_from_file_location('tlamatini_model_settings', shared)
            module = _model_import.module_from_spec(spec)
            spec.loader.exec_module(module)
            return module.resolve_agent_models('pptxer', config, agent_file=__file__)
    return config


def get_python_command() -> list:
    python_home = get_user_python_home()
    if python_home:
        exe = os.path.join(python_home, 'python.exe' if os.name == 'nt' else 'python')
        if os.path.isfile(exe):
            return [exe]
    if getattr(sys, 'frozen', False):
        return ['python']
    return [sys.executable]


def get_user_python_home() -> str:
    if getattr(sys, 'frozen', False):
        _carried = os.path.join(os.path.dirname(sys.executable), 'python')
        _exe = os.path.join(_carried, 'python.exe' if os.name == 'nt' else 'python')
        if os.path.isfile(_exe):
            return _carried
    if not sys.platform.startswith('win'):
        return os.environ.get('PYTHON_HOME', '')
    try:
        import winreg
        with winreg.OpenKey(winreg.HKEY_CURRENT_USER, r'Environment') as key:
            value, _ = winreg.QueryValueEx(key, 'PYTHON_HOME')
            return str(value) if value else ''
    except (FileNotFoundError, OSError):
        return ''


def get_agent_env() -> dict:
    env = os.environ.copy()
    if sys.platform.startswith('win'):
        try:
            import ctypes
            if hasattr(ctypes.windll.kernel32, 'SetDllDirectoryW'):
                ctypes.windll.kernel32.SetDllDirectoryW(None)
        except Exception:
            pass
    if getattr(sys, 'frozen', False) and hasattr(sys, '_MEIPASS'):
        meipass = getattr(sys, '_MEIPASS')
        if meipass:
            path_parts = env.get('PATH', '').split(os.pathsep)
            path_parts = [p for p in path_parts
                          if os.path.normpath(p) != os.path.normpath(meipass)]
            env['PATH'] = os.pathsep.join(path_parts)
    python_home = get_user_python_home()
    if not python_home:
        return env
    env['PYTHON_HOME'] = python_home
    scripts_dir = os.path.join(python_home, 'Scripts')
    current_path = env.get('PATH', '')
    env['PATH'] = f"{python_home};{scripts_dir};{current_path}"
    return env


def get_pool_path() -> str:
    current_dir = os.path.dirname(os.path.abspath(__file__))
    parent = os.path.dirname(current_dir)
    grandparent = os.path.dirname(parent)
    if os.path.basename(grandparent) == 'pools':
        return parent
    if os.path.basename(parent) == 'pools':
        return parent
    return os.path.join(os.path.dirname(current_dir), 'pools')


def get_agent_directory(agent_name: str) -> str:
    return os.path.join(get_pool_path(), agent_name)


def get_agent_script_path(agent_name: str) -> str:
    agent_dir = get_agent_directory(agent_name)
    if os.path.exists(os.path.join(agent_dir, f"{agent_name}.py")):
        return os.path.join(agent_dir, f"{agent_name}.py")
    parts = agent_name.rsplit('_', 1)
    if len(parts) == 2 and parts[1].isdigit():
        base = parts[0]
        if os.path.exists(os.path.join(agent_dir, f"{base}.py")):
            return os.path.join(agent_dir, f"{base}.py")
    return os.path.join(agent_dir, f"{agent_name}.py")


def is_agent_running(agent_name: str) -> bool:
    agent_dir = get_agent_directory(agent_name)
    pid_path = os.path.join(agent_dir, "agent.pid")
    if not os.path.exists(pid_path):
        return False
    try:
        with open(pid_path, "r") as f:
            pid = int(f.read().strip())
    except (ValueError, OSError):
        return False
    try:
        import psutil
        if not psutil.pid_exists(pid):
            return False
        proc = psutil.Process(pid)
        if proc.status() == psutil.STATUS_ZOMBIE:
            return False
        return True
    except Exception:
        try:
            os.kill(pid, 0)
            return True
        except OSError:
            return False


def wait_for_agents_to_stop(agent_names: list):
    if not agent_names:
        return
    waited = 0.0
    poll_interval = 0.5
    while True:
        still_running = [name for name in agent_names if is_agent_running(name)]
        if not still_running:
            return
        if waited >= 10.0:
            logging.error(
                f"❌ WAITING FOR AGENTS TO STOP: {still_running} still running "
                f"after {int(waited)}s. Will keep waiting..."
            )
            waited = 0.0
        time.sleep(poll_interval)
        waited += poll_interval


def start_agent(agent_name: str) -> bool:
    agent_dir = get_agent_directory(agent_name)
    script_path = get_agent_script_path(agent_name)
    if not os.path.exists(script_path):
        logging.error(f"❌ Agent script not found: {script_path}")
        return False
    try:
        cmd = get_python_command() + [script_path]
        logging.info(f"   Command: {cmd}")
        process = subprocess.Popen(
            cmd,
            cwd=agent_dir,
            env=get_agent_env(),
            creationflags=subprocess.CREATE_NO_WINDOW if sys.platform == 'win32' else 0
        )
        try:
            pid_path = os.path.join(agent_dir, "agent.pid")
            with open(pid_path, "w") as f:
                f.write(str(process.pid))
        except Exception as pid_err:
            logging.error(f"⚠️ Failed to write PID file for target {agent_name}: {pid_err}")
        logging.info(f"✅ Started agent '{agent_name}' with PID: {process.pid}")
        return True
    except Exception as e:
        logging.error(f"❌ Failed to start agent '{agent_name}': {e}")
        return False


PID_FILE = "agent.pid"


def write_pid_file():
    try:
        with open(PID_FILE, "w") as f:
            f.write(str(os.getpid()))
    except Exception as e:
        logging.error(f"❌ Failed to write PID file: {e}")


def remove_pid_file():
    for _attempt in range(5):
        try:
            if os.path.exists(PID_FILE):
                os.remove(PID_FILE)
            return
        except PermissionError:
            time.sleep(0.1)
        except Exception as e:
            logging.error(f"❌ Failed to remove PID file: {e}")
            return


# ========================================
# CONFIG COERCION
# ========================================

def _cfg(config: dict, key: str, default=""):
    value = config.get(key, default)
    return default if value is None else value


def _as_int(raw, default: int) -> int:
    """Coerce anything to an int, never raising.

    ⚠️ The wrapped chat-agent parser can deliver a value like
    "1920 pixels on the long edge" when the LLM is chatty. Extracting the
    leading number and falling back to the default is what keeps a verbose
    model from crashing the agent (the Recorder incident).
    """
    if isinstance(raw, bool):
        return default
    if isinstance(raw, (int, float)):
        return int(raw)
    try:
        import re as _re
        match = _re.search(r"-?\d+", str(raw))
        return int(match.group(0)) if match else default
    except Exception:
        return default


def _as_bool(raw, default: bool) -> bool:
    if isinstance(raw, bool):
        return raw
    if raw is None:
        return default
    text = str(raw).strip().lower()
    if text in ("true", "yes", "on", "1", "si", "sí"):
        return True
    if text in ("false", "no", "off", "0"):
        return False
    return default


def _as_list(raw) -> list:
    if raw is None:
        return []
    if isinstance(raw, (list, tuple)):
        return [str(v).strip() for v in raw if str(v).strip()]
    text = str(raw).strip()
    if not text:
        return []
    if text.startswith("["):
        try:
            parsed = json.loads(text)
            if isinstance(parsed, list):
                return [str(v).strip() for v in parsed if str(v).strip()]
        except (ValueError, TypeError):
            pass
    parts = [p.strip() for p in text.replace(";", ",").split(",")]
    return [p for p in parts if p]


# ========================================
# THE SIBLING MODULES (fail-open group import)
# ========================================

_MODULES = {}
_MODULE_ERROR = ""

try:
    import pptxer_audit
    import pptxer_build
    import pptxer_color
    import pptxer_diagram          # noqa: F401  (used via pptxer_build)
    import pptxer_docmodel
    import pptxer_draw             # noqa: F401  (used via pptxer_build)
    import pptxer_fonts
    import pptxer_layout
    import pptxer_media
    import pptxer_nuance
    import pptxer_render
    import pptxer_theme

    _MODULES = {
        "audit": pptxer_audit, "build": pptxer_build, "color": pptxer_color,
        "docmodel": pptxer_docmodel, "fonts": pptxer_fonts,
        "layout": pptxer_layout, "media": pptxer_media,
        "nuance": pptxer_nuance, "render": pptxer_render, "theme": pptxer_theme,
    }
except Exception as exc:                                          # noqa: BLE001
    _MODULE_ERROR = str(exc)


# ========================================
# OUTPUT LOCATION
# ========================================

def _documents_dir() -> str:
    """The Windows Documents KNOWN FOLDER, which is not always ~/Documents.

    A user with OneDrive redirection, or a localized Windows ("Documentos"),
    has a Documents folder that a hardcoded path misses entirely — the deck
    would land somewhere they never look.
    """
    if sys.platform.startswith("win"):
        try:
            import ctypes
            import ctypes.wintypes
            folder_id = ctypes.create_unicode_buffer(
                "{FDD39AD0-238F-46AF-ADB4-6C85480369C7}")
            guid = ctypes.wintypes.GUID if hasattr(ctypes.wintypes, "GUID") else None
            del folder_id, guid
        except Exception:                                         # noqa: BLE001
            pass
        try:
            import ctypes
            from ctypes import windll, wintypes

            class GUID(ctypes.Structure):
                _fields_ = [("Data1", wintypes.DWORD), ("Data2", wintypes.WORD),
                            ("Data3", wintypes.WORD), ("Data4", ctypes.c_byte * 8)]

            documents = GUID(0xFDD39AD0, 0x238F, 0x46AF,
                             (ctypes.c_byte * 8)(0xAD, 0xB4, 0x6C, 0x85,
                                                 0x48, 0x03, 0x69, 0xC7))
            path_ptr = ctypes.c_wchar_p()
            if windll.shell32.SHGetKnownFolderPath(
                    ctypes.byref(documents), 0, None,
                    ctypes.byref(path_ptr)) == 0:
                value = path_ptr.value
                windll.ole32.CoTaskMemFree(path_ptr)
                if value:
                    return value
        except Exception:                                         # noqa: BLE001
            pass
    return os.path.join(os.path.expanduser("~"), "Documents")


def _dir_accepts_new_files(path: str) -> bool:
    try:
        os.makedirs(path, exist_ok=True)
        probe = os.path.join(path, f".pptxer_probe_{os.getpid()}")
        with open(probe, "w", encoding="utf-8") as fh:
            fh.write("ok")
        os.remove(probe)
        return True
    except Exception:                                             # noqa: BLE001
        return False


def _default_output_dir(config: dict) -> str:
    explicit = str(_cfg(config, "output_dir", "")).strip()
    if explicit:
        return os.path.abspath(os.path.expandvars(os.path.expanduser(explicit)))
    return os.path.join(_documents_dir(), "TlamatiniPPTX")


def _temp_root(subdir="") -> str:
    root = (os.environ.get("TLAMATINI_TEMP") or "").strip()
    if not root:
        root = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Temp")
    target = os.path.join(root, "PPTXer", subdir) if subdir else os.path.join(root, "PPTXer")
    try:
        os.makedirs(target, exist_ok=True)
    except Exception:                                             # noqa: BLE001
        pass
    return target


def _resolve_output_path(config: dict) -> str:
    out_dir = _default_output_dir(config)
    # basename() so `filename` can never escape output_dir — a path-traversal
    # guard that costs nothing and closes the hole permanently.
    name = os.path.basename(str(_cfg(config, "filename", "")).strip())
    if not name:
        stamp = time.strftime("%Y%m%d_%H%M%S")
        millis = int((time.time() % 1) * 1000)
        name = f"pptxer_{stamp}_{millis:03d}.pptx"
    if not name.lower().endswith(".pptx"):
        name += ".pptx"

    path = os.path.join(out_dir, name)
    if _as_bool(_cfg(config, "overwrite", False), False):
        return path

    base, ext = os.path.splitext(path)
    counter = 2
    while os.path.exists(path) and counter < 1000:
        path = f"{base}_{counter}{ext}"
        counter += 1
    return path


# ========================================
# FAIL-SAFE PREFLIGHT
# ========================================

def _preflight(action: str, config: dict, content: str) -> dict:
    """REFUSE rather than write a wrong deck. Never crashes.

    Returns {ok, blockers, warnings, stage}. A blocker means PPTXer declines
    with `status: refused` — which is a routable outcome a downstream Forker
    can branch on, not an error.
    """
    blockers = []
    warnings = []

    if _MODULE_ERROR:
        blockers.append(
            f"PPTXer's design modules could not be imported ({_MODULE_ERROR}). "
            f"The agent directory is incomplete — every pptxer_*.py sibling must "
            f"sit beside pptxer.py.")
        return {"ok": False, "blockers": blockers, "warnings": warnings,
                "stage": "modules"}

    try:
        import pptx                                              # noqa: F401
    except Exception as exc:                                      # noqa: BLE001
        blockers.append(
            f"python-pptx is not available ({exc}). PPTXer cannot write a "
            f"presentation without it; install it into the carried Python.")

    try:
        from PIL import Image                                    # noqa: F401
    except Exception as exc:                                      # noqa: BLE001
        warnings.append(
            f"Pillow is unavailable ({exc}) — text cannot be measured with real "
            f"font metrics and no artwork can be generated. The deck will still "
            f"be written, with conservative estimates.")

    if action in ("create", "outline"):
        if not content.strip():
            blockers.append(
                "no content was supplied — set input_text, or input_file to a "
                ".md / .txt / .json file.")

    if action in ("render", "audit", "info"):
        target = str(_cfg(config, "pptx_path", "")).strip()
        if not target:
            blockers.append(f"action '{action}' needs pptx_path.")
        elif not os.path.isfile(target):
            blockers.append(f"the deck was not found: {target}")

    if action == "create":
        out_dir = _default_output_dir(config)
        if not _dir_accepts_new_files(out_dir):
            blockers.append(
                f"the output directory cannot be written: {out_dir}")

    size = str(_cfg(config, "slide_size", "16:9")).strip().lower()
    if size and size not in _MODULES["layout"].SLIDE_SIZES:
        warnings.append(
            f"slide_size {size!r} is unknown; falling back to 16:9. Valid: "
            f"{', '.join(sorted(_MODULES['layout'].SLIDE_SIZES))}")

    nuance = str(_cfg(config, "nuance", "")).strip()
    if nuance and not _MODULES["nuance"].resolve_nuance(nuance):
        warnings.append(
            f"nuance {nuance!r} is not in the catalog; PPTXer will detect the "
            f"treatment from the content instead.")

    for key in ("predominant_color", "accent_color", "text_color",
                "background_color", "heading_color"):
        raw = str(_cfg(config, key, "")).strip()
        if raw and _MODULES["color"].parse_color(raw) is None:
            warnings.append(
                f"{key}={raw!r} could not be parsed as a colour and was ignored.")

    return {"ok": not blockers, "blockers": blockers, "warnings": warnings,
            "stage": "preflight"}


def _format_preflight(pf: dict) -> str:
    lines = []
    for blocker in pf.get("blockers", []):
        lines.append(f"  ⛔ {blocker}")
    for warning in pf.get("warnings", []):
        lines.append(f"  ⚠️  {warning}")
    return "\n".join(lines)


# ========================================
# CONTENT RESOLUTION
# ========================================

def _resolve_content(config: dict) -> tuple:
    """(text, source_kind). input_text wins; then input_file."""
    text = str(_cfg(config, "input_text", "") or "")
    if text.strip():
        return (text, "input_text")

    path = str(_cfg(config, "input_file", "")).strip()
    if not path:
        return ("", "none")

    resolved = os.path.abspath(os.path.expandvars(os.path.expanduser(path)))
    if not os.path.isfile(resolved):
        logging.error(f"❌ input_file not found: {resolved}")
        return ("", "missing")

    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            with open(resolved, "r", encoding=encoding) as fh:
                return (fh.read(), f"file:{os.path.basename(resolved)}")
        except (UnicodeDecodeError, UnicodeError):
            continue
        except Exception as exc:                                  # noqa: BLE001
            logging.error(f"❌ input_file could not be read: {exc}")
            return ("", "unreadable")
    return ("", "unreadable")


def _collect_media(deck, config: dict) -> tuple:
    """Fetch and prepare every asset the deck references. → (index, failures).

    Failures are KEPT, not dropped: one dead URL must not silently shrink a
    gallery, and the agent must be able to say which asset died and why.
    """
    media = _MODULES["media"]
    index = {}
    failures = []

    max_px = _as_int(_cfg(config, "max_image_px", 1920), 1920)
    timeout = _as_int(_cfg(config, "fetch_timeout", 25), 25)
    max_bytes = _as_int(_cfg(config, "max_media_bytes", 67108864), 67108864)
    allow_private = _as_bool(_cfg(config, "allow_private_hosts", False), False)
    dest = _temp_root("media")

    sources = []
    for slide in deck.slides:
        sources.extend(slide.images)
    sources.extend(_as_list(_cfg(config, "images", [])))

    seen = set()
    for src in sources:
        if not src or src in seen:
            continue
        seen.add(src)
        asset = media.prepare_image(src, max_px=max_px, dest_dir=dest,
                                    allow_private=allow_private,
                                    max_bytes=max_bytes, timeout=timeout)
        index[src] = asset
        if asset.ok:
            note = f" ({asset.note})" if asset.note else ""
            logging.info(f"🖼️  image ready: {asset.width}x{asset.height} "
                         f"{asset.bytes // 1024}KB {src[:70]}{note}")
        else:
            failures.append(f"image {src[:90]} — {asset.error}")
            logging.error(f"❌ image unavailable: {src[:70]} — {asset.error}")

    for slide in deck.slides:
        if not slide.video:
            continue
        if slide.video in index:
            continue
        asset = media.prepare_video(slide.video, dest_dir=dest,
                                    allow_private=allow_private,
                                    max_bytes=max_bytes, timeout=timeout)
        index[slide.video] = asset
        if asset.ok:
            logging.info(f"🎬 video ready: {asset.width}x{asset.height} "
                         f"{asset.duration:.1f}s {slide.video[:60]}")
        else:
            failures.append(f"video {slide.video[:90]} — {asset.error}")
            logging.error(f"❌ video unavailable: {asset.error}")

    hero_video = str(_cfg(config, "video", "")).strip()
    if hero_video and hero_video not in index:
        asset = media.prepare_video(hero_video, dest_dir=dest,
                                    allow_private=allow_private,
                                    max_bytes=max_bytes, timeout=timeout)
        index[hero_video] = asset
        if asset.ok and deck.slides:
            target = next((s for s in deck.slides if not s.video), None)
            if target is not None:
                target.video = hero_video
                target.kind = "media"
        elif not asset.ok:
            failures.append(f"video {hero_video[:90]} — {asset.error}")

    hero_audio = str(_cfg(config, "audio", "")).strip()
    if hero_audio:
        asset = media.prepare_audio(hero_audio, dest_dir=dest,
                                    allow_private=allow_private,
                                    max_bytes=max_bytes, timeout=timeout)
        index[hero_audio] = asset
        if not asset.ok:
            failures.append(f"audio {hero_audio[:90]} — {asset.error}")

    return (index, failures)


# ========================================
# PARAMETRIZER SECTION
# ========================================

def _emit_section(fields: dict, body: str) -> None:
    """ONE atomic logging.info call. Concurrent writes interleave otherwise."""
    header = "\n".join(f"{key}: {value}" for key, value in fields.items())
    payload = f"INI_SECTION_PPTXER<<<\n{header}\n"
    if body:
        payload += f"\n{body}\n"
    payload += ">>>END_SECTION_PPTXER"
    logging.info(payload)


# ========================================
# ACTIONS
# ========================================

def _action_fonts(config: dict) -> dict:
    inv = _MODULES["fonts"].get_inventory()
    lines = [f"{inv.family_count} font families are installed on this machine.",
             f"Scanned: {', '.join(inv.scanned_dirs)}", ""]
    for category, names in _MODULES["fonts"].FONT_CATEGORIES.items():
        present = inv.installed_from(names)
        lines.append(f"{category} ({len(present)} of {len(names)} available):")
        lines.append("  " + (", ".join(present) if present else "— none —"))
    lines.append("")
    lines.append("Pairings that resolve on this machine:")
    for name in _MODULES["fonts"].PAIRINGS:
        pair = _MODULES["fonts"].resolve_pairing(name, inv)
        mark = "exact " if not pair["substitutions"] else "subbed"
        lines.append(f"  [{mark}] {name:<18} {pair['display'].family} / "
                     f"{pair['body'].family} / {pair['mono'].family}")
    return {"ok": True, "status": "listed", "body": "\n".join(lines),
            "font_families": inv.family_count}


def _action_validate(config: dict) -> dict:
    lines = ["PPTXer backend and renderer report:", ""]
    for label, module in (("python-pptx", "pptx"), ("Pillow", "PIL"),
                          ("PyMuPDF", "fitz"), ("OpenCV", "cv2"),
                          ("NumPy", "numpy")):
        try:
            __import__(module)
            lines.append(f"  OK   {label}")
        except Exception as exc:                                  # noqa: BLE001
            lines.append(f"  MISS {label} — {exc}")

    lines.append("")
    lines.append("Renderers (how PPTXer can look at its own slides):")
    for tier, info in _MODULES["render"].probe_renderers().items():
        mark = "OK  " if info["available"] else "MISS"
        lines.append(f"  {mark} {tier}: {info['detail']}")
        lines.append(f"       {info['quality']}")

    inv = _MODULES["fonts"].get_inventory()
    lines.append("")
    lines.append(f"Fonts: {inv.family_count} families available for measurement.")
    return {"ok": True, "status": "validated", "body": "\n".join(lines),
            "font_families": inv.family_count}


def _action_info(config: dict) -> dict:
    path = str(_cfg(config, "pptx_path", "")).strip()
    try:
        from pptx import Presentation
        prs = Presentation(path)
        slides = len(prs.slides)
        shapes = sum(len(s.shapes) for s in prs.slides)
        pictures = 0
        tables = 0
        charts = 0
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.shape_type == 13:
                    pictures += 1
                if getattr(shape, "has_table", False):
                    tables += 1
                if getattr(shape, "has_chart", False):
                    charts += 1
        size_emu = (int(prs.slide_width or 0), int(prs.slide_height or 0))
        body = (f"{os.path.basename(path)}\n"
                f"  slides     : {slides}\n"
                f"  shapes     : {shapes}\n"
                f"  pictures   : {pictures}\n"
                f"  tables     : {tables}\n"
                f"  charts     : {charts}\n"
                f"  slide size : {size_emu[0]}x{size_emu[1]} EMU "
                f"({size_emu[0] / 914400:.2f}x{size_emu[1] / 914400:.2f} in)\n"
                f"  bytes      : {os.path.getsize(path):,}")
        return {"ok": True, "status": "listed", "body": body,
                "slide_count": slides, "output_path": path,
                "bytes": os.path.getsize(path)}
    except Exception as exc:                                      # noqa: BLE001
        return {"ok": False, "status": "error",
                "body": f"the deck could not be read: {exc}"}


def _action_render_or_audit(config: dict, action: str) -> dict:
    path = str(_cfg(config, "pptx_path", "")).strip()
    render_result = None
    if action == "render":
        render_result = _MODULES["render"].render_slides(
            path, out_dir=_temp_root("shots"),
            width=_as_int(_cfg(config, "render_width", 1600), 1600),
            prefer=str(_cfg(config, "render", "auto")).strip().lower(),
            timeout=_as_int(_cfg(config, "render_timeout", 240), 240))
    report = _MODULES["audit"].audit_pptx(
        path, render_result=render_result,
        margin_pt=_as_int(_cfg(config, "margin_pt", 48), 48),
        run_pixels=bool(render_result))
    return {
        "ok": True,
        "status": "ok" if report.layout_clean else "findings",
        "body": report.summary(),
        "layout_clean": report.layout_clean,
        "overlaps": report.overlap_count if hasattr(report, "overlap_count")
                    else len(report.overlaps),
        "slide_count": report.slides_measured,
        "output_path": path,
        "rendered": (render_result or {}).get("count", 0),
        "report": report,
        "render_result": render_result,
    }


def _action_outline(config: dict, content: str) -> dict:
    fmt = str(_cfg(config, "input_format", "auto")).strip().lower()
    deck = _MODULES["docmodel"].parse_content(content, max_bullets=6, fmt=fmt)
    title = str(_cfg(config, "title", "")).strip()
    if title:
        deck.title = title

    verdict = _MODULES["nuance"].classify(
        deck.plain_text(), title=deck.title,
        requested=str(_cfg(config, "nuance", "")).strip(),
        image_count=deck.image_count)

    lines = [f"Treatment : {verdict.label} (confidence {verdict.confidence:.2f})",
             f"Reasoning : {verdict.reasoning()}", "",
             f"{len(deck)} slides, {deck.word_count} words:"]
    for i, slide in enumerate(deck.slides, start=1):
        detail = slide.title or slide.body[:44] or slide.quote[:44] or "—"
        lines.append(f"  {i:>2}. [{slide.kind:<14}] {detail[:62]}")
    return {"ok": True, "status": "listed", "body": "\n".join(lines),
            "slide_count": len(deck), "nuance": verdict.nuance,
            "nuance_confidence": round(verdict.confidence, 2)}


def _action_create(config: dict, content: str) -> dict:
    docmodel = _MODULES["docmodel"]
    nuance_mod = _MODULES["nuance"]
    theme_mod = _MODULES["theme"]
    build_mod = _MODULES["build"]
    render_mod = _MODULES["render"]
    audit_mod = _MODULES["audit"]
    fonts_mod = _MODULES["fonts"]

    inventory = fonts_mod.get_inventory()
    logging.info(f"🔤 {inventory.family_count} font families available for measurement")

    # ---- 1. content -> deck model -----------------------------------------
    fmt = str(_cfg(config, "input_format", "auto")).strip().lower()
    probe_nuance = nuance_mod.classify(
        content[:6000], title=str(_cfg(config, "title", "")),
        requested=str(_cfg(config, "nuance", "")).strip())
    density = str(_cfg(config, "density", "")).strip() or \
        nuance_mod.NUANCES.get(probe_nuance.nuance, {}).get("density", "medium")
    max_bullets = {"minimal": 3, "low": 5, "medium": 6, "high": 8}.get(density, 6)

    deck = docmodel.parse_content(content, max_bullets=max_bullets, fmt=fmt)
    for key, attr in (("title", "title"), ("subtitle", "subtitle"),
                      ("author", "author")):
        value = str(_cfg(config, key, "")).strip()
        if value:
            setattr(deck, attr, value)
    if not deck.title:
        deck.title = "Presentation"

    logging.info(f"📑 parsed {len(deck)} slides from {deck.source_format} "
                 f"({deck.word_count} words): {deck.kinds()}")

    # ---- 2. cover + closing ------------------------------------------------
    kicker = str(_cfg(config, "kicker", "")).strip()
    cover = docmodel.Slide(kind="title_slide", title=deck.title,
                           subtitle=deck.subtitle, kicker=kicker,
                           notes=str(_cfg(config, "notes", "")))
    deck.slides.insert(0, cover)

    # ---- 3. what KIND of deck is this? ------------------------------------
    verdict = nuance_mod.classify(
        deck.plain_text(), title=deck.title,
        requested=str(_cfg(config, "nuance", "")).strip(),
        image_count=deck.image_count, slide_hint=len(deck))
    logging.info(f"🎨 treatment: {verdict.label} "
                 f"(confidence {verdict.confidence:.2f}, {verdict.source})")
    logging.info(f"   {verdict.reasoning()}")

    # ---- 4. the design system ---------------------------------------------
    theme = theme_mod.build_theme(verdict, config, inventory)
    for line in theme.describe().splitlines():
        logging.info(f"   {line}")

    # ---- 5. media ----------------------------------------------------------
    media_index, media_failures = _collect_media(deck, config)

    # ---- 6. build ----------------------------------------------------------
    output_path = _resolve_output_path(config)
    logging.info(f"🏗️  building {len(deck)} slides -> {output_path}")
    build = build_mod.build_deck(deck, theme, output_path, config,
                                 media_index=media_index, inventory=inventory)
    build.media_failures.extend(media_failures)

    if not build.ok:
        return {"ok": False, "status": "error",
                "body": f"the presentation could not be written: {build.error}",
                "slide_count": build.slide_count}

    logging.info(f"✅ wrote {build.slide_count} slides, {build.shape_count} shapes, "
                 f"{os.path.getsize(output_path):,} bytes")

    # ---- 7. LOOK AT IT -----------------------------------------------------
    prefer = str(_cfg(config, "render", "auto")).strip().lower()
    shots_dir = (os.path.join(os.path.dirname(output_path),
                              os.path.splitext(os.path.basename(output_path))[0] + "_slides")
                 if _as_bool(_cfg(config, "save_slide_images", True), True)
                 else _temp_root("shots"))
    size_key = str(_cfg(config, "slide_size", "16:9")).strip().lower()
    size_spec = _MODULES["layout"].SLIDE_SIZES.get(
        size_key, _MODULES["layout"].SLIDE_SIZES["16:9"])
    render_result = render_mod.render_slides(
        output_path, out_dir=shots_dir,
        width=_as_int(_cfg(config, "render_width", 1600), 1600),
        prefer=prefer, slide_specs=build.slide_specs,
        slide_size=(size_spec[0], size_spec[1]),
        timeout=_as_int(_cfg(config, "render_timeout", 240), 240))

    if render_result.get("ok"):
        logging.info(f"👁️  rendered {render_result['count']} slides via "
                     f"{render_result['tier']} "
                     f"(ground truth: {render_result['ground_truth']})")
    else:
        for attempt in render_result.get("attempts", []):
            logging.info(f"   renderer {attempt['tier']}: "
                         f"available={attempt['available']} "
                         f"{attempt.get('error', '')[:90]}")

    # ---- 8. MEASURE IT -----------------------------------------------------
    report = None
    if _as_bool(_cfg(config, "layout_audit", True), True):
        report = audit_mod.audit_pptx(
            output_path, render_result=render_result,
            margin_pt=_as_int(_cfg(config, "margin_pt", 48), 48),
            text_boxes=build.text_boxes, inventory=inventory)
        for line in report.summary().splitlines():
            logging.info(f"   {line}")

    body_parts = [theme.describe(), ""]
    if report is not None:
        body_parts.append(report.summary())
    if build.warnings:
        body_parts.append("")
        body_parts.append("Build notes:")
        body_parts.extend(f"  · {w}" for w in build.warnings[:12])
    if build.media_failures:
        body_parts.append("")
        body_parts.append("Media that could not be used:")
        body_parts.extend(f"  · {m}" for m in build.media_failures[:12])
    if theme.meta.get("font_substitutions"):
        body_parts.append("")
        body_parts.append("Font substitutions on this machine:")
        body_parts.extend(f"  · {s}" for s in theme.meta["font_substitutions"])

    layout_clean = report.layout_clean if report is not None else None
    status = "created"
    if report is not None and not report.layout_clean:
        # A deck with measured defects is NOT reported as a clean success.
        status = "created_with_findings"

    return {
        "ok": True,
        "status": status,
        "body": "\n".join(body_parts),
        "output_path": output_path,
        "slide_count": build.slide_count,
        "shape_count": build.shape_count,
        "images_embedded": build.images_embedded,
        "videos_embedded": build.videos_embedded,
        "tables": build.tables,
        "charts": build.charts,
        "diagrams": build.diagrams,
        "generated_art": build.generated_art,
        "nuance": verdict.nuance,
        "nuance_confidence": round(verdict.confidence, 2),
        "theme": theme,
        "layout_clean": layout_clean,
        "report": report,
        "render_result": render_result,
        "bytes": os.path.getsize(output_path),
    }


# ========================================
# MAIN
# ========================================

def main():
    config = load_config()
    write_pid_file()
    started = time.time()
    outcome = {}

    try:
        if _IS_REANIMATED:
            logging.info(f"🔄 {CURRENT_DIR_NAME} REANIMATED (resuming from pause)")
            logging.info("=" * 60)

        target_agents = _as_list(config.get('target_agents', []))
        action = str(_cfg(config, "action", "create")).strip().lower() or "create"

        logging.info("🖼️ PPTXER AGENT STARTED")
        logging.info(f"🎯 Targets: {target_agents}")
        logging.info(f"⚙️  Action: {action}")
        logging.info("=" * 60)

        content, source_kind = _resolve_content(config)
        if action in ("create", "outline"):
            logging.info(f"📥 content from {source_kind} ({len(content)} chars)")

        # ---- fail-safe preflight -----------------------------------------
        if _as_bool(_cfg(config, "preflight", True), True):
            pf = _preflight(action, config, content)
            report = _format_preflight(pf)
            if report:
                logging.info("🛡️  Preflight:")
                for line in report.splitlines():
                    logging.info(line)
            if not pf["ok"]:
                logging.error("⛔ REFUSING — preflight found a blocker. "
                              "PPTXer declines rather than write a wrong deck.")
                outcome = {
                    "ok": False, "status": "refused", "stage": pf["stage"],
                    "body": ("PPTXer refused rather than produce a wrong or empty "
                             "presentation.\n\n" + report),
                }
                raise _Refused()
        else:
            logging.info("🛡️  Preflight DISABLED by configuration")

        # ---- dispatch -----------------------------------------------------
        if action == "fonts":
            outcome = _action_fonts(config)
        elif action == "validate":
            outcome = _action_validate(config)
        elif action == "info":
            outcome = _action_info(config)
        elif action in ("render", "audit"):
            outcome = _action_render_or_audit(config, action)
        elif action == "outline":
            outcome = _action_outline(config, content)
        elif action == "create":
            outcome = _action_create(config, content)
        else:
            outcome = {
                "ok": False, "status": "refused", "stage": "action",
                "body": (f"unknown action {action!r}. Valid: create, outline, "
                         f"render, audit, info, fonts, validate."),
            }

    except _Refused:
        pass
    except Exception as exc:                                      # noqa: BLE001
        logging.error(f"❌ PPTXer failed: {exc}")
        import traceback
        for line in traceback.format_exc().splitlines()[-12:]:
            logging.error(f"   {line}")
        outcome = {"ok": False, "status": "error",
                   "body": f"PPTXer failed unexpectedly: {exc}"}

    # ---- emit the section, ALWAYS -----------------------------------------
    try:
        theme = outcome.get("theme")
        report = outcome.get("report")
        render_result = outcome.get("render_result") or {}

        fields = {
            "action": str(_cfg(config, "action", "create")).strip().lower(),
            "status": outcome.get("status", "error"),
            "success": bool(outcome.get("ok", False)),
            "output_path": outcome.get("output_path", ""),
            "output_dir": os.path.dirname(outcome.get("output_path", "")) or "",
            "filename": os.path.basename(outcome.get("output_path", "")) or "",
            "slide_count": outcome.get("slide_count", 0),
            "shape_count": outcome.get("shape_count", 0),
            "images_embedded": outcome.get("images_embedded", 0),
            "videos_embedded": outcome.get("videos_embedded", 0),
            "tables": outcome.get("tables", 0),
            "charts": outcome.get("charts", 0),
            "diagrams": outcome.get("diagrams", 0),
            "generated_art": outcome.get("generated_art", 0),
            "nuance": outcome.get("nuance", ""),
            "nuance_confidence": outcome.get("nuance_confidence", 0),
            "palette": (theme.color("accent").hex if theme is not None else ""),
            "predominant_color": (theme.meta.get("seed", "")
                                  if theme is not None else ""),
            "font_display": (theme.fonts["display"].family
                             if theme is not None else ""),
            "font_body": (theme.fonts["body"].family if theme is not None else ""),
            "font_families": outcome.get(
                "font_families",
                # Report the real inventory for EVERY action, not only `fonts`
                # and `validate` — a deck that says "0 font families" reads as
                # a broken machine when it simply was not asked.
                (_MODULES["fonts"].get_inventory().family_count
                 if _MODULES else 0)),
            "shapes_measured": getattr(report, "shapes_measured", 0),
            "render_tier": render_result.get("tier", ""),
            "slides_rendered": render_result.get("count", 0),
            "ground_truth": bool(render_result.get("ground_truth", False)),
            "layout_clean": (report.layout_clean if report is not None
                             else outcome.get("layout_clean", "")),
            "overlaps": (len(report.overlaps) if report is not None else 0),
            "text_overflows": (len(report.text_overflows) if report is not None else 0),
            "bytes": outcome.get("bytes", 0),
            "elapsed_seconds": round(time.time() - started, 2),
            "stage": outcome.get("stage", "done"),
        }
        _emit_section(fields, outcome.get("body", ""))
    except Exception as exc:                                      # noqa: BLE001
        logging.error(f"⚠️ the result section could not be emitted: {exc}")

    # ---- ALWAYS trigger downstream ----------------------------------------
    # Success, failure OR fail-safe refusal — so a downstream Forker can branch
    # on {status} / {layout_clean} / {slide_count} instead of the flow dying.
    try:
        target_agents = _as_list(config.get('target_agents', []))
        if target_agents:
            wait_for_agents_to_stop(target_agents)
            for target in target_agents:
                start_agent(target)
        else:
            logging.info("ℹ️ No downstream agents configured.")
    except Exception as exc:                                      # noqa: BLE001
        logging.error(f"❌ downstream trigger failed: {exc}")

    logging.info(f"🏁 PPTXer agent finished. Status: {outcome.get('status', 'error')}")

    try:
        time.sleep(0.4)
    finally:
        remove_pid_file()
    sys.exit(0)


class _Refused(Exception):
    """Internal control flow for a fail-safe refusal. Never escapes main()."""


if __name__ == "__main__":
    main()
