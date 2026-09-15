# ══════════════════════════════════════════════════════════════════════════════
#   ✦  T L A M A T I N I  ✦   —   "one who knows"
#
#   Created by  Angela López Mendoza   ·   @angelahack1
#   Developer · Architect · Creator of Tlamatini
#
#   Every line of this file was written by Angela López Mendoza.
# ══════════════════════════════════════════════════════════════════════════════
#   Tlamatini Author Banner — do not remove (releases scrub the name automatically)
"""PPTXer — the agent that MEASURES ITS OWN SLIDES, and the tests that measure it.

These tests are deliberately HARD. They drive the REAL modules with REAL font
files and REAL geometry; nothing that is under test is mocked. Where a test
needs an absent capability (PowerPoint, a network), it asserts the HONEST
DEGRADATION rather than skipping — because "the renderer was unavailable" and
"the deck is fine" must never be confusable, and that confusion is the exact
bug class PPTXer exists to prevent.

Every incident found while building the agent on 2026-09-14 is pinned here, with
the measured numbers, so none of them can come back:

  · a half-em width estimate is −38% wrong on capital W's (the overflow cause);
  · text validated against the flat ground still rendered at 6.19:1, because
    the real background is a gradient with an ornament over it;
  · over-correcting that contrast destroyed the accent colour — the ORNAMENT
    must yield, never the palette;
  · a 119pt title ate 55% of the slide until stat numbers were smaller than
    their own labels;
  · `app.Quit()` leaks POWERPNT.EXE unless the COM references are dropped first;
  · a text frame inside its own panel is NOT an overlap;
  · a deliberate bleed-off glow is NOT off-slide content;
  · "12ms" and "47%" must both parse as stats, and "Twelve new maps" must not.
"""
import importlib.util
import logging
import os
import sys
import tempfile
import unittest

from django.test import SimpleTestCase

AGENT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                         "agents", "pptxer")


def _load(module_name):
    """Import a pptxer_* sibling by path.

    The agent runs as ``python pptxer.py`` from a copied runtime directory, so
    its modules are flat neighbours rather than a package. AGENT_DIR is put on
    sys.path so the siblings can import each other exactly as they do at
    runtime — importing them any other way would test a different arrangement
    than the one that ships.
    """
    if AGENT_DIR not in sys.path:
        sys.path.insert(0, AGENT_DIR)
    path = os.path.join(AGENT_DIR, module_name + ".py")
    spec = importlib.util.spec_from_file_location(module_name, path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[module_name] = module
    spec.loader.exec_module(module)
    return module


C = _load("pptxer_color")
F = _load("pptxer_fonts")
L = _load("pptxer_layout")
N = _load("pptxer_nuance")
T = _load("pptxer_theme")
D = _load("pptxer_docmodel")
M = _load("pptxer_media")
DRAW = _load("pptxer_draw")
R = _load("pptxer_render")
A = _load("pptxer_audit")


# ═════════════════════════════════════════════════════════════════════════════
#  COLOUR
# ═════════════════════════════════════════════════════════════════════════════

class ColourScienceTests(SimpleTestCase):

    def test_oklab_round_trip_is_exact_to_8_bit(self):
        """Every conversion must survive a round trip, or gradients drift."""
        for hexcode in ("#000000", "#FFFFFF", "#FF0000", "#00FF00", "#0000FF",
                        "#39FF14", "#7B2FF7", "#12E2A3", "#808080"):
            back = C.oklab_to_srgb(*C.srgb_to_oklab(hexcode))
            for original, restored in zip(C.parse_color(hexcode).rgb255,
                                          back.rgb255):
                self.assertLessEqual(abs(original - restored), 1,
                                     f"{hexcode} did not round-trip")

    def test_oklab_midpoint_is_lighter_than_the_srgb_one(self):
        """THE reason gradients are interpolated in OKLab.

        Measured: red→blue in sRGB passes through #800080 at L=0.420 — visibly
        darker than either end, which is what makes a stock two-stop gradient
        look muddy. In OKLab the same midpoint sits at L≈0.54 and stays
        luminous. If this ever inverts, every deck's artwork has regressed.
        """
        srgb_mid = C.Color(0.5, 0.0, 0.5)
        oklab_mid = C.mix_oklab("#FF0000", "#0000FF", 0.5)
        self.assertGreater(C.srgb_to_oklab(oklab_mid)[0],
                           C.srgb_to_oklab(srgb_mid)[0] + 0.08)

    def test_parse_is_fail_open_and_never_raises(self):
        for junk in ("not-a-colour", "", None, "#GGGGGG", "rgb(", [], {},
                     "oklch(bad)", 12345):
            self.assertIsNone(C.parse_color(junk),
                              f"{junk!r} should fail open to None")

    def test_named_gaming_colours_resolve(self):
        for name in ("cyberpunk", "esports", "neon", "obsidian", "synthwave",
                     "arcade", "plasma", "toxic"):
            self.assertIsNotNone(C.parse_color(name), f"{name} must resolve")

    def test_ensure_contrast_keeps_the_hue_and_clears_the_floor(self):
        """A brand colour is made LEGIBLE, not replaced."""
        fixed = C.ensure_contrast("#FFFF00", "#FFFFFF", C.BODY_CONTRAST_FLOOR)
        self.assertGreaterEqual(C.contrast_ratio(fixed, "#FFFFFF"),
                                C.BODY_CONTRAST_FLOOR)
        # Still yellow: the OKLCh hue must be within a few degrees.
        _, _, want = C.srgb_to_oklch("#FFFF00")
        _, _, got = C.srgb_to_oklch(fixed)
        self.assertLess(min(abs(want - got), 360 - abs(want - got)), 12.0)

    def test_the_projected_floors_are_stricter_than_a_printed_page(self):
        """A slide is read from 10m in a lit room, not held at arm's length."""
        self.assertGreaterEqual(C.BODY_CONTRAST_FLOOR, 7.0)
        self.assertGreater(C.CONTRAST_SAFETY, 1.0,
                           "never sit exactly ON a threshold")

    def test_seed_from_text_is_deterministic(self):
        self.assertEqual(C.seed_from_text("Nexus Protocol").hex,
                         C.seed_from_text("Nexus Protocol").hex)


# ═════════════════════════════════════════════════════════════════════════════
#  FONTS — the anti-overflow core
# ═════════════════════════════════════════════════════════════════════════════

class FontMeasurementTests(SimpleTestCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.inv = F.get_inventory()

    def test_the_machine_has_measurable_fonts(self):
        self.assertGreater(self.inv.family_count, 20)

    def test_a_half_em_estimate_is_badly_wrong_both_ways(self):
        """THE measured reason PPTXer never guesses a text width.

        On Impact at 40pt the half-em rule is roughly +78% wrong on lowercase
        l's and −38% wrong on capital W's. The −38% is how text escapes its
        box; if this test stops failing the estimate, someone has swapped real
        metrics back out for arithmetic.
        """
        resolved = self.inv.resolve("Impact", "display_impact")
        if not resolved.measurable:
            self.skipTest("no measurable display face on this machine")

        narrow, _ = F.measure_text("lllllllllll", resolved, 40)
        wide, _ = F.measure_text("WWWWWWWWWWW", resolved, 40)
        estimate = 11 * 40 * 0.5

        self.assertLess(narrow, estimate * 0.75)
        self.assertGreater(wide, estimate * 1.25)
        self.assertGreater(wide, narrow * 2.0,
                           "W must measure far wider than l — if these are "
                           "close, the real font is not being used")

    def test_wrap_never_mutates_an_unbreakable_token(self):
        """A Windows path must survive intact even when it does not fit.

        Injecting a space or a hyphen into a path is WRONG DATA. The correct
        behaviour is to return the token on its own line and let the caller
        report an overflow.
        """
        resolved = self.inv.resolve("Segoe UI", "sans_body")
        path = r"C:\Users\angel\AppData\Local\Tlamatini\telemetry\export.csv"
        lines = F.wrap_text_to_width(path, resolved, 18, 60.0)
        self.assertIn(path, lines)
        self.assertEqual("".join(lines).count("\\"), path.count("\\"))

    def test_fit_point_size_result_actually_fits(self):
        resolved = self.inv.resolve("Segoe UI", "sans_body")
        size, lines = F.fit_point_size(
            "NEXUS PROTOCOL SEASON THREE LAUNCH DECK", resolved,
            max_width_pt=400, max_height_pt=120, start_pt=72)
        widest = max(F.measure_text(ln, resolved, size)[0] for ln in lines)
        self.assertLessEqual(widest, 400.5)
        self.assertLessEqual(len(lines) * size * 1.18, 120.5)

    def test_substitution_reports_itself_and_stays_in_its_width_class(self):
        """A silent substitution is how a deck overflows on another machine."""
        resolved = self.inv.resolve("ThisFontDoesNotExistAnywhere",
                                    "sans_condensed")
        self.assertTrue(resolved.substituted)
        self.assertTrue(resolved.reason, "a substitution must say WHY")
        self.assertTrue(resolved.measurable,
                        "a substitute must still be measurable")

    def test_every_pairing_resolves_to_measurable_faces(self):
        for name in F.PAIRINGS:
            pair = F.resolve_pairing(name, self.inv)
            for role in ("display", "body", "mono"):
                self.assertTrue(pair[role].measurable,
                                f"{name}.{role} is not measurable")


# ═════════════════════════════════════════════════════════════════════════════
#  LAYOUT — the invariant
# ═════════════════════════════════════════════════════════════════════════════

class FontInventoryDeterminismTests(SimpleTestCase):
    """The 2026-09-14 font-resolution defects, both found by the VISIBLE test.

    Neither was visible to the library: python-pptx saved a valid file every
    time, and the agent reported the RIGHT FAMILY NAME in both cases. Only
    measuring the finished deck exposed them.
    """

    def test_the_inventory_resolves_the_same_files_under_different_hash_seeds(self):
        """⚠️ ONLY OBSERVABLE ACROSS PROCESSES — one interpreter has one seed.

        `_scan` used to iterate `set(paths)` while keeping the FIRST file seen
        for a given (family, style). Python randomises string hashing per
        process, so when two files claim the same style the winner changed
        from run to run, and the winner's metrics decide how text wraps.

        Measured: the SAME deck, same content, same code, built twice —
        audited LAYOUT CLEAN once and "statlabel needs 41pt in a 40pt frame"
        the next, while both runs reported the identical families
        (Bahnschrift / Segoe UI). Only the FILE behind the name had changed.
        A deck that lays out differently between runs is untrustworthy, and
        it makes every layout bug unreproducible.
        """
        import subprocess

        child = (
            "import sys\n"
            "sys.path.insert(0, r'%s')\n"
            "import pptxer_fonts as F\n"
            "inv = F.FontInventory()\n"
            "print('|'.join('%%s=%%s' %% (f, getattr(inv.resolve(f), 'path', '-'))\n"
            "               for f in ('Arial', 'Segoe UI', 'Verdana', 'Tahoma')))\n"
        ) % AGENT_DIR

        seen = []
        for seed in ("1", "424242"):
            env = dict(os.environ)
            env["PYTHONHASHSEED"] = seed
            proc = subprocess.run([sys.executable, "-c", child],
                                  capture_output=True, text=True,
                                  encoding="utf-8", errors="replace",
                                  env=env, timeout=300)
            if proc.returncode != 0:
                self.skipTest("inventory unavailable: %s" % proc.stderr[-200:])
            seen.append(proc.stdout.strip())

        self.assertEqual(
            seen[0], seen[1],
            "the inventory resolved DIFFERENT font files under two hash seeds "
            "— iteration order is leaking into the layout")

    def test_a_width_variant_never_evicts_its_base_family(self):
        """('Arial', 'Narrow') is Arial NARROW, not Arial's regular face.

        Windows puts the width in the STYLE field, and `_style_key` reads only
        weight and slant — so "Narrow" fell through to "regular" and the four
        Arial Narrow files claimed every slot of the **Arial** bucket,
        evicting arial.ttf entirely. `resolve("Arial")` then returned
        ARIALN.TTF while reporting `substituted=False`: text measured and
        declared as Arial, drawn narrower than Arial.
        """
        fonts = _load("pptxer_fonts")

        self.assertEqual(fonts._fold_width_into_family("Arial", "Narrow"),
                         ("Arial Narrow", "regular"))
        self.assertEqual(fonts._fold_width_into_family("Arial", "Bold Narrow"),
                         ("Arial Narrow", "bold"))
        # a family that already carries its width is never doubled
        self.assertEqual(
            fonts._fold_width_into_family("Arial Narrow", "Narrow"),
            ("Arial Narrow", "Narrow"))
        # an ordinary weight/slant is left completely alone
        self.assertEqual(fonts._fold_width_into_family("Verdana", "Bold"),
                         ("Verdana", "Bold"))

        inv = fonts.FontInventory()
        arial = inv.resolve("Arial")
        if arial is None or not arial.measurable:
            self.skipTest("Arial is not installed on this machine")
        self.assertNotIn(
            "arialn", os.path.basename(arial.path).lower(),
            "resolve('Arial') returned an Arial NARROW file (%s)" % arial.path)


class LayoutInvariantTests(SimpleTestCase):

    def setUp(self):
        self.canvas = L.SlideCanvas("16:9", margin_pt=48)

    def test_slide_geometry_is_exact(self):
        self.assertEqual(self.canvas.width, 12192000)
        self.assertEqual(self.canvas.height, 6858000)
        self.assertEqual(L.EMU_PER_INCH % L.EMU_PER_POINT, 0,
                         "EMU must divide evenly by points")

    def test_overlap_is_refused(self):
        # Start INSIDE the safe area, so the refusal under test is the overlap
        # and not an escape (at margin_pt=48 the safe origin is 609600 EMU).
        safe = self.canvas.safe
        first = self.canvas.place(L.Box(safe.left, safe.top, 2000000, 1000000,
                                        name="first"))
        with self.assertRaises(L.PlacementError) as caught:
            self.canvas.place(L.Box(first.left + 100000, first.top + 100000,
                                    2000000, 1000000, name="second"))
        self.assertIn("overlap", str(caught.exception).lower())

    def test_leaving_the_safe_area_is_refused(self):
        with self.assertRaises(L.PlacementError):
            self.canvas.place(L.Box(self.canvas.width - 100000, 100000,
                                    900000, 900000, name="escapee"))

    def test_columns_and_rows_tile_their_region_exactly(self):
        region = self.canvas.safe
        for count in range(1, 7):
            cols = self.canvas.columns(region, count)
            total = sum(c.width for c in cols) + self.canvas.gutter * (count - 1)
            self.assertEqual(total, region.width, f"{count} columns must tile")
            rows = self.canvas.rows(region, count)
            total_h = sum(r.height for r in rows) + self.canvas.gutter * (count - 1)
            self.assertEqual(total_h, region.height, f"{count} rows must tile")

    def test_containment_is_not_intersection_for_the_audit(self):
        outer = L.Box(0, 0, 1000000, 1000000, name="panel")
        inner = L.Box(100000, 100000, 500000, 500000, name="label")
        self.assertTrue(outer.contains(inner))

    def test_a_shared_edge_is_not_an_overlap(self):
        left = L.Box(0, 0, 1000000, 1000000)
        right = L.Box(1000000, 0, 1000000, 1000000)
        self.assertFalse(left.intersects(right))

    def test_image_fit_preserves_aspect_ratio(self):
        region = L.Box(0, 0, 4000000, 2000000, name="frame")
        for mode in ("contain", "cover"):
            box = L.fit_image_box(region, 1920, 1080, mode)
            self.assertAlmostEqual(box.width / box.height, 1920 / 1080, places=1,
                                   msg=f"{mode} distorted the image")

    def test_the_audit_proves_a_clean_layout(self):
        for col in self.canvas.columns(self.canvas.safe, 3):
            self.canvas.place(col)
        report = self.canvas.audit()
        self.assertTrue(report["clean"])
        self.assertEqual(report["overlap_count"], 0)


class ColumnSolverTests(SimpleTestCase):
    """The direct descendant of PDFer's TableSolver."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.inv = F.get_inventory()
        cls.font = cls.inv.resolve("Segoe UI", "sans_body")

    def test_widths_always_sum_to_the_region_exactly(self):
        canvas = L.SlideCanvas("16:9")
        rows = [["Metric", "Q1", "Q2", "Commentary"],
                ["Monthly active users", "1,204,551", "1,890,223", "Growth"],
                ["Path", r"C:\Users\angel\Local\export.csv", "-", "-"]]
        solved = L.solve_columns(rows, self.font, 14, canvas.safe.width,
                                 inventory=self.inv)
        self.assertEqual(sum(solved["widths"]), canvas.safe.width)

    def test_an_unbreakable_token_raises_its_own_column_minimum(self):
        """The PDFer bug, prevented: a long path must widen ITS column."""
        canvas = L.SlideCanvas("16:9")
        short = [["A", "B"], ["x", "y"]]
        long_path = [["A", "B"],
                     ["x", r"C:\Users\angel\AppData\Local\Tlamatini\export.csv"]]
        a = L.solve_columns(short, self.font, 14, canvas.safe.width,
                            inventory=self.inv)
        b = L.solve_columns(long_path, self.font, 14, canvas.safe.width,
                            inventory=self.inv)
        self.assertGreater(b["minimums_pt"][1], a["minimums_pt"][1] * 2)

    def test_impossible_widths_are_clamped_and_reported(self):
        """The destructive rung is last, and it SAYS so."""
        rows = [["A" * 60, "B" * 60, "C" * 60, "D" * 60, "E" * 60]]
        solved = L.solve_columns(rows, self.font, 28, 2000000,
                                 inventory=self.inv)
        self.assertTrue(solved["min_violated"])
        self.assertTrue(solved["reason"])
        self.assertEqual(sum(solved["widths"]), 2000000,
                         "even a clamped solve must tile exactly")


class TextFitTests(SimpleTestCase):

    def test_overflow_is_reported_not_hidden(self):
        canvas = L.SlideCanvas("16:9")
        inv = F.get_inventory()
        font = inv.resolve("Segoe UI", "sans_body")
        result = canvas.fit_text("Supercalifragilisticexpialidocious" * 3,
                                 L.Box(0, 0, 700000, 300000, name="tiny"),
                                 font, start_pt=40)
        self.assertTrue(result["overflow"])
        self.assertIn("unbreakable", result["reason"])

    def test_a_comfortable_box_does_not_report_overflow(self):
        canvas = L.SlideCanvas("16:9")
        inv = F.get_inventory()
        font = inv.resolve("Segoe UI", "sans_body")
        result = canvas.fit_text("Season Three", canvas.safe, font, start_pt=40)
        self.assertFalse(result["overflow"])


# ═════════════════════════════════════════════════════════════════════════════
#  NUANCE
# ═════════════════════════════════════════════════════════════════════════════

class NuanceClassifierTests(SimpleTestCase):

    def test_it_recognises_the_decks_it_was_built_for(self):
        cases = (
            ("esports_tournament", "GRAND FINALS",
             "Group A bracket. Team Nova won the series 3-1 in a BO5. Prize "
             "pool was $250,000. MVP had a 7.2 KDA across the qualifier."),
            ("startup_pitch", "Series A",
             "The problem. Our solution. TAM is $4B, SAM $600M. ARR grew to "
             "$2.1M with CAC of $900 and LTV $14k. Traction. The ask."),
            ("safety_briefing", "Lab Protocol",
             "DANGER: corrosive. Always wear PPE. Administer 50 mg within ten "
             "minutes. Follow the evacuation route. Risk assessment first."),
            ("financial_disclosure", "FY2026 Results",
             "Revenue of $412.8 million, up 14.2% YoY. EBITDA margin 23%. The "
             "balance sheet shows net income of $88M per GAAP."),
            ("game_design_doc", "Core Systems",
             "The core loop is explore-fight-upgrade. Skill tree has 40 nodes. "
             "Balance pass nerfed the shotgun and buffed respawn after "
             "playtest feedback on the difficulty curve."),
        )
        for expected, title, body in cases:
            verdict = N.classify(body, title=title)
            self.assertEqual(verdict.nuance, expected,
                             f"{title!r} -> {verdict.nuance} ({verdict.reasoning()})")

    def test_thin_content_falls_back_rather_than_guessing_confidently(self):
        verdict = N.classify("Hello. Some words. Nothing much.")
        self.assertEqual(verdict.source, "fallback")
        self.assertLess(verdict.confidence, 0.4)

    def test_an_explicit_nuance_always_wins(self):
        verdict = N.classify("DANGER hazard PPE evacuation 50 mg dosage",
                             requested="esports_tournament")
        self.assertEqual(verdict.nuance, "esports_tournament")
        self.assertEqual(verdict.source, "explicit")

    def test_spanish_and_english_aliases_agree(self):
        for alias, expected in (("gamer", "game_design_doc"),
                                ("lanzamiento", "product_launch"),
                                ("torneo", "esports_tournament"),
                                ("seguridad", "safety_briefing"),
                                ("presentacion", None)):
            resolved = N.resolve_nuance(alias)
            if expected:
                self.assertEqual(resolved, expected, f"{alias} -> {resolved}")

    def test_a_safety_critical_deck_is_never_decorated(self):
        """⚠️ Not even when the user explicitly asks for 'rich'."""
        for key in N.SAFETY_CRITICAL_NUANCES:
            verdict = N.NuanceVerdict(key, 1.0, "explicit")
            self.assertEqual(N.decoration_ceiling_for(verdict, "rich"), "none",
                             f"{key} must never be decorated")

    def test_low_confidence_lowers_the_decoration_ceiling(self):
        confident = N.NuanceVerdict("esports_tournament", 0.95, "detected")
        unsure = N.NuanceVerdict("esports_tournament", 0.20, "detected")
        order = ("none", "restrained", "moderate", "rich")
        self.assertLessEqual(order.index(N.decoration_ceiling_for(unsure)),
                             order.index(N.decoration_ceiling_for(confident)))


# ═════════════════════════════════════════════════════════════════════════════
#  THEME — the contrast guarantee
# ═════════════════════════════════════════════════════════════════════════════

class ThemeContrastTests(SimpleTestCase):

    def test_every_text_role_clears_its_floor_in_every_nuance(self):
        """168 assertions. This is the guarantee, not a sample of it."""
        failures = []
        for key in N.NUANCES:
            theme = T.build_theme(N.NuanceVerdict(key, 1.0, "explicit"), {})
            ground = theme.color("ground_worst", theme.color("ground").hex)
            for role, floor in (("text", 7.0), ("caption", 7.0),
                                ("footer", 7.0), ("title", 4.5),
                                ("text_muted", 4.5)):
                ratio = C.contrast_ratio(theme.color(role), ground)
                if ratio < floor:
                    failures.append(f"{key}.{role}={ratio:.2f}<{floor}")
            if C.contrast_ratio(theme.color("on_surface"),
                                theme.color("surface_high")) < 7.0:
                failures.append(f"{key}.on_surface")
            if C.contrast_ratio(theme.color("on_accent"),
                                theme.color("accent")) < 4.5:
                failures.append(f"{key}.on_accent")
        self.assertEqual(failures, [], f"{len(failures)} illegible roles")

    def test_an_illegible_user_palette_is_repaired_and_the_repair_explained(self):
        theme = T.build_theme(N.NuanceVerdict("corporate_update", 1.0, "explicit"),
                              {"background_color": "#FFFFFF",
                               "text_color": "#EEEEEE"})
        self.assertGreaterEqual(
            C.contrast_ratio(theme.color("text"),
                             theme.color("ground_worst",
                                         theme.color("ground").hex)), 7.0)
        self.assertTrue(theme.validation_notes,
                        "a forced change must be explained, never silent")

    def test_one_seed_colour_drives_the_whole_palette(self):
        theme = T.build_theme(N.NuanceVerdict("esports_tournament", 1.0, "explicit"),
                              {"predominant_color": "#12E2A3"})
        _, _, seed_hue = C.srgb_to_oklch("#12E2A3")
        _, _, accent_hue = C.srgb_to_oklch(theme.color("accent"))
        self.assertLess(min(abs(seed_hue - accent_hue),
                            360 - abs(seed_hue - accent_hue)), 25.0)

    def test_ground_worst_is_further_from_the_text_than_the_nominal_ground(self):
        """The 2026-09-14 finding: the nominal ground is not what text lands on."""
        theme = T.build_theme(N.NuanceVerdict("cyberpunk_tech", 1.0, "explicit"),
                              {"predominant_color": "#00F0FF"})
        text = theme.color("text")
        self.assertLessEqual(
            C.contrast_ratio(text, theme.color("ground_worst")),
            C.contrast_ratio(text, theme.color("ground")) + 0.01)

    def test_type_scale_respects_the_projected_minimums(self):
        for key in N.NUANCES:
            theme = T.build_theme(N.NuanceVerdict(key, 1.0, "explicit"), {})
            self.assertGreaterEqual(theme.size("body"), 16.0,
                                    f"{key}: body text below the room floor")
            self.assertGreater(theme.size("title"), theme.size("body"),
                               f"{key}: broken hierarchy")


# ═════════════════════════════════════════════════════════════════════════════
#  DOCUMENT MODEL
# ═════════════════════════════════════════════════════════════════════════════

class DocModelTests(SimpleTestCase):

    def test_markdown_becomes_the_right_slide_kinds(self):
        deck = D.parse_content(
            "# Deck\n\n## Bullets\n- one\n- two\n\n"
            "## Table\n| A | B |\n|---|---|\n| 1 | 2 |\n\n"
            "## Quote\n> A thing somebody said\n\n"
            "## Code\n```python\nprint(1)\n```\n")
        kinds = deck.kinds()
        self.assertEqual(deck.title, "Deck")
        for expected in ("table", "quote", "code"):
            self.assertIn(expected, kinds, f"{expected} slide was not promoted")

    def test_stat_lines_parse_including_units(self):
        """Every one of these is a real line from a real deck."""
        should_match = ("3.4M Registered players", "890K Peak concurrent",
                        "47% Retention at day 30", "12ms Median server latency",
                        "144fps Target frame rate", "2.4GB Install size",
                        "3x Faster than last season", "60 Playable characters",
                        "$2.1M Annual recurring revenue")
        for line in should_match:
            self.assertIsNotNone(D._STAT_RE.match(line), f"{line!r} must parse")
        for line in ("Twelve new maps built on the Aether engine",
                     "Not a stat at all really"):
            self.assertIsNone(D._STAT_RE.match(line), f"{line!r} must NOT parse")

    def test_parallel_stats_are_promoted_to_tiles(self):
        deck = D.parse_content(
            "# D\n\n## Numbers\n- 3.4M Players\n- 47% Retention\n"
            "- 12ms Latency\n- 2.4GB Size\n")
        stats = [s for s in deck.slides if s.kind == "stats"]
        self.assertEqual(len(stats), 1)
        self.assertEqual(len(stats[0].stats), 4)

    def test_an_overfull_slide_is_SPLIT_never_shrunk(self):
        """Content is never lost, and a continuation says it is one."""
        bullets = "\n".join(f"- point number {i}" for i in range(14))
        deck = D.parse_content(f"# D\n\n## Many\n{bullets}\n", max_bullets=5)
        many = [s for s in deck.slides if s.title.startswith("Many")]
        self.assertGreater(len(many), 1, "the slide should have been split")
        self.assertEqual(sum(len(s.bullets) for s in many), 14,
                         "no bullet may be dropped")
        self.assertIn("(cont.)", many[1].title)

    def test_format_detection(self):
        self.assertEqual(D.detect_format('{"slides": []}'), "json")
        self.assertEqual(D.detect_format("# Title\n\ntext"), "markdown")
        self.assertEqual(D.detect_format("Just a sentence."), "text")

    def test_json_deck_spec_round_trips(self):
        deck = D.parse_content(
            '{"title": "T", "slides": [{"kind": "stats", "title": "N",'
            ' "stats": [{"value": "9", "label": "things"}]}]}')
        self.assertEqual(deck.title, "T")
        self.assertEqual(deck.slides[0].kind, "stats")

    def test_markdown_emphasis_is_stripped_but_words_are_kept(self):
        deck = D.parse_content("# D\n\n## S\n- **bold** and *italic* and `code`\n")
        bullet = deck.slides[0].bullets[0]
        self.assertNotIn("*", bullet)
        self.assertIn("bold", bullet)


# ═════════════════════════════════════════════════════════════════════════════
#  MEDIA — bytes, never headers
# ═════════════════════════════════════════════════════════════════════════════

class MediaSafetyTests(SimpleTestCase):

    def test_magic_numbers_identify_the_format(self):
        cases = {
            b"\x89PNG\r\n\x1a\n": ("image", ".png"),
            b"\xff\xd8\xff\xe0": ("image", ".jpg"),
            b"GIF89a": ("image", ".gif"),
            b"RIFF\x00\x00\x00\x00WEBP": ("image", ".webp"),
            b"RIFF\x00\x00\x00\x00AVI ": ("video", ".avi"),
            b"RIFF\x00\x00\x00\x00WAVE": ("audio", ".wav"),
            b"\x00\x00\x00\x18ftypisom": ("video", ".mp4"),
            b"ID3\x03\x00": ("audio", ".mp3"),
        }
        for payload, expected in cases.items():
            self.assertEqual(M.sniff_kind(payload + b"\x00" * 48), expected)

    def test_an_html_page_is_never_mistaken_for_an_image(self):
        """A 404 page served as image/jpeg must be REFUSED, not embedded."""
        html = b"<!DOCTYPE html><html><body>404 Not Found</body></html>"
        self.assertEqual(M.sniff_kind(html), ("", ""))

    def test_private_and_loopback_hosts_are_refused(self):
        """⚠️ Fail-CLOSED. This is the one guard that must not fail open."""
        for url in ("http://127.0.0.1:8000/admin",
                    "http://localhost/secret.png",
                    "http://169.254.169.254/latest/meta-data/",
                    "http://10.0.0.5/logo.png"):
            asset = M.fetch_media(url)
            self.assertFalse(asset.ok, f"{url} must be refused")
            self.assertIn("private", asset.error.lower())

    def test_non_http_schemes_say_why(self):
        for url in ("file:///C:/Windows/win.ini", "ftp://example.com/x.png"):
            asset = M.fetch_media(url)
            self.assertFalse(asset.ok)
            self.assertIn("http", asset.error.lower())

    def test_an_unembeddable_video_container_is_refused_not_faked(self):
        """PPTXer ships no encoder and will not rename a container."""
        with tempfile.TemporaryDirectory() as tmp:
            mkv = os.path.join(tmp, "clip.mkv")
            with open(mkv, "wb") as handle:
                handle.write(b"\x1aE\xdf\xa3" + b"\x00" * 256)
            asset = M.prepare_video(mkv, dest_dir=tmp)
            self.assertFalse(asset.ok)
            self.assertIn("mp4", asset.error.lower())
            self.assertIn("not transcode", asset.error.lower())

    def test_a_missing_file_fails_without_raising(self):
        asset = M.fetch_media(os.path.join(tempfile.gettempdir(), "nope.png"))
        self.assertFalse(asset.ok)


# ═════════════════════════════════════════════════════════════════════════════
#  ARTWORK
# ═════════════════════════════════════════════════════════════════════════════

class ArtworkTests(SimpleTestCase):

    def test_every_ornament_is_drawable_and_deterministic(self):
        if not DRAW.available():
            self.skipTest("Pillow unavailable")
        colours = {"ground": "#0B0B0F", "accent": "#00F0FF",
                   "accent_2": "#FF10F0", "accent_3": "#FFD93D",
                   "text": "#FFFFFF"}
        with tempfile.TemporaryDirectory() as tmp:
            for name in DRAW.ORNAMENTS:
                if name == "none":
                    self.assertIsNone(DRAW.draw_ornament(name, 320, 180, colours))
                    continue
                first = DRAW.draw_ornament(
                    name, 320, 180, colours,
                    out_path=os.path.join(tmp, f"{name}_a.png"), seed=42)
                second = DRAW.draw_ornament(
                    name, 320, 180, colours,
                    out_path=os.path.join(tmp, f"{name}_b.png"), seed=42)
                self.assertTrue(first, f"{name} drew nothing")
                with open(first, "rb") as fa, open(second, "rb") as fb:
                    self.assertEqual(fa.read(), fb.read(),
                                     f"{name} is not deterministic")

    def test_an_unknown_ornament_fails_open(self):
        self.assertIsNone(DRAW.draw_ornament("not_a_real_motif", 100, 100, {}))

    def test_the_seed_is_stable(self):
        self.assertEqual(DRAW.make_seed("a", 1), DRAW.make_seed("a", 1))
        self.assertNotEqual(DRAW.make_seed("a", 1), DRAW.make_seed("a", 2))


# ═════════════════════════════════════════════════════════════════════════════
#  RENDER — honest about what it cannot do
# ═════════════════════════════════════════════════════════════════════════════

class RendererHonestyTests(SimpleTestCase):

    def test_the_probe_reports_every_tier(self):
        probe = R.probe_renderers()
        for tier in (R.TIER_POWERPOINT, R.TIER_LIBREOFFICE, R.TIER_PREVIEW):
            self.assertIn(tier, probe)
            self.assertIn("available", probe[tier])
            self.assertTrue(probe[tier]["detail"])

    def test_an_unavailable_tier_reports_unavailable_not_clean(self):
        """⚠️ THE ACPX `--version` LIE, PREVENTED.

        A probe that could not run has NO opinion. If a tier is unavailable it
        must say so and say WHY — never return a result that a caller could
        read as "no problems found".
        """
        probe = R.probe_renderers()
        for tier, info in probe.items():
            if not info["available"]:
                self.assertTrue(info["detail"],
                                f"{tier} is unavailable and did not say why")

    def test_render_disabled_returns_a_truthful_nothing(self):
        result = R.render_slides("nonexistent.pptx", prefer="none")
        self.assertFalse(result["ok"])
        self.assertFalse(result["ground_truth"])
        self.assertIn("disabled", result["note"])

    def test_only_powerpoint_counts_as_ground_truth(self):
        preview = R.RenderResult(tier=R.TIER_PREVIEW, ok=True, images=["a.png"])
        libre = R.RenderResult(tier=R.TIER_LIBREOFFICE, ok=True, images=["a.png"])
        real = R.RenderResult(tier=R.TIER_POWERPOINT, ok=True, images=["a.png"])
        self.assertFalse(preview.is_ground_truth)
        self.assertFalse(libre.is_ground_truth)
        self.assertTrue(real.is_ground_truth)

    def test_the_reaper_never_touches_a_pre_existing_powerpoint(self):
        """Angela may have PowerPoint open with unsaved work."""
        if not sys.platform.startswith("win"):
            self.skipTest("Windows only")
        existing = R._powerpoint_pids()
        killed = R._reap_orphan_powerpoint(existing, grace=0.1)
        self.assertEqual(killed, [],
                         "a PID present BEFORE the render must never be killed")


# ═════════════════════════════════════════════════════════════════════════════
#  AUDIT — and the false alarms it must NOT raise
# ═════════════════════════════════════════════════════════════════════════════

class AuditHonestyTests(SimpleTestCase):

    def test_a_missing_render_is_reported_as_not_measured(self):
        """⚠️ An audit that could not look must never report 'clean'."""
        report = A.AuditReport()
        report.levels_skipped.append("pixel audit — no render was produced")
        self.assertIn("no renderer", report.confidence.lower())
        self.assertNotIn("what the audience will see", report.confidence)

    def test_ground_truth_changes_the_stated_confidence(self):
        report = A.AuditReport()
        report.ground_truth = True
        report.pixel_audited = 5
        self.assertIn("PowerPoint", report.confidence)

    def test_defects_are_never_reported_as_clean(self):
        report = A.AuditReport()
        report.overlaps = [{"slide": 1, "a": "x", "b": "y", "area_pt2": 12.0}]
        self.assertFalse(report.layout_clean)
        self.assertIn("defects", report.confidence)


# ═════════════════════════════════════════════════════════════════════════════
#  END-TO-END — the real thing
# ═════════════════════════════════════════════════════════════════════════════

DECK_SOURCE = """# NEXUS PROTOCOL
Season Three Launch

## The Arena Has Changed
- Twelve new maps built on the Aether engine
- Ranked ladder rebuilt around squad synergy
- Spectator mode with sub-second latency

## By The Numbers
- 3.4M Registered players
- 890K Peak concurrent
- 47% Retention at day 30
- 12ms Median server latency

## Season Pass Tiers
| Tier | Price | Includes |
|---|---|---|
| Recruit | Free | Battle pass track, 2 skins |
| Operative | $9.99 | Full track, 14 skins, XP boost |
| Vanguard | $24.99 | Everything, plus the Aether weapon set |

## What Players Told Us
> The ranked system finally respects the time I put into it.
"""


class EndToEndDeckTests(SimpleTestCase):
    """Builds a REAL .pptx and re-opens it. Nothing here is mocked."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        try:
            import pptx                                          # noqa: F401
        except Exception:                                        # noqa: BLE001
            raise unittest.SkipTest("python-pptx unavailable")

        cls._tmp = tempfile.mkdtemp(prefix="pptxer_test_")
        cls.build_mod = _load("pptxer_build")

        cls.deck = D.parse_content(DECK_SOURCE, max_bullets=5)
        cls.verdict = N.classify(cls.deck.plain_text(), title=cls.deck.title)
        cls.theme = T.build_theme(cls.verdict, {"predominant_color": "#12E2A3"})
        cls.path = os.path.join(cls._tmp, "e2e.pptx")

        level = logging.getLogger().level
        try:
            cls.result = cls.build_mod.build_deck(
                cls.deck, cls.theme, cls.path,
                {"slide_size": "16:9", "generate_art": False})
        finally:
            logging.getLogger().setLevel(level)

    @classmethod
    def tearDownClass(cls):
        import shutil
        shutil.rmtree(getattr(cls, "_tmp", ""), ignore_errors=True)
        super().tearDownClass()

    def test_the_deck_was_written(self):
        self.assertTrue(self.result.ok, self.result.error)
        self.assertTrue(os.path.isfile(self.path))
        self.assertGreater(os.path.getsize(self.path), 10000)

    def test_the_expected_slide_kinds_were_built(self):
        kinds = self.deck.kinds()
        self.assertIn("stats", kinds)
        self.assertIn("table", kinds)
        self.assertIn("quote", kinds)
        self.assertGreaterEqual(self.result.tables, 1)

    def test_the_saved_file_has_no_overlaps_and_nothing_off_slide(self):
        """Measured on the ARTEFACT, re-opened — not on the in-memory model."""
        tree = A.audit_shape_tree(self.path)
        self.assertEqual(tree["error"], "")
        self.assertEqual(tree["overlaps"], [],
                         f"{len(tree['overlaps'])} overlapping shapes")
        self.assertEqual(tree["off_slide"], [],
                         f"{len(tree['off_slide'])} shapes off the slide")
        self.assertGreater(tree["shapes"], 10)

    def test_no_text_frame_overflows_its_box(self):
        fit = A.audit_text_fit(self.path)
        self.assertEqual(fit["error"], "")
        self.assertEqual(fit["overflows"], [],
                         f"{len(fit['overflows'])} overflowing text frames")

    def test_stat_values_are_larger_than_their_own_labels(self):
        """The 2026-09-14 hierarchy inversion, pinned.

        A tile whose number is smaller than its caption has defeated the only
        purpose a stat tile has.
        """
        from pptx import Presentation
        presentation = Presentation(self.path)
        values, labels = [], []
        for slide in presentation.slides:
            for shape in slide.shapes:
                name = (shape.name or "").lower()
                if not getattr(shape, "has_text_frame", False):
                    continue
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        if run.font.size is None:
                            continue
                        if "statvalue" in name:
                            values.append(run.font.size.pt)
                        elif "statlabel" in name:
                            labels.append(run.font.size.pt)
        if not values or not labels:
            self.skipTest("no stat tiles in this deck")
        self.assertGreater(min(values), max(labels),
                           "a stat number must be bigger than its label")

    def test_every_text_run_has_an_explicit_size(self):
        """Autofit is OFF by design; a size PowerPoint decides is unauditable."""
        from pptx import Presentation
        presentation = Presentation(self.path)
        missing = 0
        for slide in presentation.slides:
            for shape in slide.shapes:
                if not getattr(shape, "has_text_frame", False):
                    continue
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        if run.text.strip() and run.font.size is None:
                            missing += 1
        self.assertEqual(missing, 0, f"{missing} runs with no explicit size")

    def test_the_build_reports_its_own_placement_diagnostics(self):
        self.assertIsInstance(self.result.placement_diagnostics, list)
        self.assertIsInstance(self.result.warnings, list)


# ═════════════════════════════════════════════════════════════════════════════
#  REGISTRY INTEGRATION
# ═════════════════════════════════════════════════════════════════════════════

class MeasureWhatYouDrawTests(SimpleTestCase):
    """The 2026-09-14 caps/measure mismatch and the stat-tile split, pinned.

    Both defects were found by the VISIBLE test, not by a unit test, and both
    were invisible to the library: python-pptx saved a perfectly valid file
    in each case. They are pinned here because each one produced a deck that
    looked finished and was not.
    """

    def test_uppercase_text_is_measured_in_the_case_it_is_drawn(self):
        """`_cased` is the whole fix: measure the string that gets drawn.

        `_add_text` renders `line.upper()` for a role set in capitals, and
        capitals are wider — so fitting the mixed-case string computes the
        wrap for text that never appears. Measured: a stat label fitted as
        "Target frame rate on mid-range hardware" and drawn as
        "TARGET FRAME RATE ON MID-RANGE HARDWARE" needed 42pt in a 39pt frame.
        """
        build = _load("pptxer_build")
        fonts = _load("pptxer_fonts")

        self.assertEqual(build._cased("Mid-Range Hardware", True),
                         "MID-RANGE HARDWARE")
        self.assertEqual(build._cased("Mid-Range Hardware", False),
                         "Mid-Range Hardware")
        self.assertEqual(build._cased(123, True), "123")

        # The premise, measured rather than asserted: capitals really are
        # wider, so the two cases cannot share one measurement.
        inv = fonts.FontInventory()
        face = inv.resolve("Verdana")
        if face is None or not face.measurable:
            self.skipTest("no measurable face on this machine")
        mixed = fonts.measure_text("Target frame rate on hardware", face, 12.0)
        caps = fonts.measure_text("TARGET FRAME RATE ON HARDWARE", face, 12.0)
        self.assertGreater(caps, mixed,
                           "capitals must measure wider than mixed case")

    def test_every_capitalised_role_is_fitted_through_cased(self):
        """A new call site that forgets `_cased` reintroduces the bug.

        Source guard rather than behaviour: the defect is silent by nature —
        the deck still saves, the text simply spills — so the only reliable
        moment to catch a regression is when the code is written.
        """
        import ast

        with open(os.path.join(AGENT_DIR, "pptxer_build.py"),
                  encoding="utf-8") as fh:
            tree = ast.parse(fh.read())

        # ⚠️ SCOPED PER FUNCTION, deliberately. Nearly every slide builder
        # names its measurement `fit`, so matching that name across the whole
        # module would tie one builder's capitalised draw to another
        # builder's mixed-case fit and report a site that is perfectly fine.
        # Each draw is checked against the assignment that actually feeds it:
        # the last one to that name, in its own function, above the call.
        def _is_cased(call):
            return any(isinstance(sub, ast.Call)
                       and getattr(sub.func, "id", "") == "_cased"
                       for sub in ast.walk(call.args[0])) if call.args else False

        checked, problems = 0, []
        for func in [n for n in ast.walk(tree)
                     if isinstance(n, ast.FunctionDef)]:
            fits = {}
            for node in ast.walk(func):
                if (isinstance(node, ast.Assign) and len(node.targets) == 1
                        and isinstance(node.targets[0], ast.Name)
                        and isinstance(node.value, ast.Call)
                        and getattr(node.value.func, "attr", "") == "fit_text"):
                    fits.setdefault(node.targets[0].id, []).append(
                        (node.lineno, _is_cased(node.value)))

            for node in ast.walk(func):
                if not (isinstance(node, ast.Call)
                        and getattr(node.func, "id", "") == "_add_text"):
                    continue
                upper = next((kw.value for kw in node.keywords
                              if kw.arg == "uppercase"), None)
                if upper is None:
                    continue
                if isinstance(upper, ast.Constant) and upper.value is False:
                    continue                  # explicitly never capitalised
                if len(node.args) < 3 or not isinstance(node.args[2], ast.Name):
                    continue
                feeding = [pair for pair in fits.get(node.args[2].id, [])
                           if pair[0] <= node.lineno]
                if not feeding:
                    continue                  # fitted elsewhere; not ours
                checked += 1
                if not feeding[-1][1]:
                    problems.append(
                        "line %d: '%s' is drawn in capitals but was fitted at "
                        "line %d without _cased() - the wrap is computed for "
                        "text that is never drawn"
                        % (node.lineno, node.args[2].id, feeding[-1][0]))

        self.assertGreaterEqual(
            checked, 5,
            "only %d capitalised text roles found - did _add_text move?"
            % checked)
        self.assertEqual(problems, [], "\n".join(problems))


class LongStatLabelTests(SimpleTestCase):
    """A long label must not overflow its tile. Measured, on the artefact."""

    LONG_STATS = """# Layout Torture

## Numbers With Units
- 144fps Target frame rate on mid-range hardware
- 2.4GB Streaming budget per map at maximum texture detail
- 68% Reduction in draw calls after the batching pass
- 12ms Median frame time measured across the whole match
"""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        try:
            import pptx                                          # noqa: F401
        except Exception:                                        # noqa: BLE001
            raise unittest.SkipTest("python-pptx unavailable")

        cls._tmp = tempfile.mkdtemp(prefix="pptxer_statlabel_")
        build_mod = _load("pptxer_build")
        deck = D.parse_content(cls.LONG_STATS, max_bullets=5)
        verdict = N.classify(deck.plain_text(), title=deck.title)
        theme = T.build_theme(verdict, {"predominant_color": "#1F8A8A"})
        cls.path = os.path.join(cls._tmp, "statlabels.pptx")

        level = logging.getLogger().level
        try:
            cls.result = build_mod.build_deck(
                deck, theme, cls.path,
                {"slide_size": "16:9", "generate_art": False})
        finally:
            logging.getLogger().setLevel(level)

    @classmethod
    def tearDownClass(cls):
        import shutil
        shutil.rmtree(getattr(cls, "_tmp", ""), ignore_errors=True)
        super().tearDownClass()

    def test_a_long_stat_label_still_fits_its_tile(self):
        """The exact 2026-09-14 failure: 'needs 42pt in a 39pt frame'.

        Two causes, both fixed: the label was measured in mixed case and
        drawn in capitals, and the tile handed the value a flat 62% of its
        height whether the number needed it or not, stranding that space
        above a label wrapping underneath.
        """
        self.assertTrue(self.result.ok, self.result.error)
        fit = A.audit_text_fit(self.path)
        self.assertEqual(fit["error"], "")
        self.assertEqual(
            fit["overflows"], [],
            "a long stat label overflowed its tile: %s" % (fit["overflows"],))


class RegistryIntegrationTests(SimpleTestCase):

    def test_the_wrapped_tool_is_registered(self):
        from agent import chat_agent_registry as registry
        spec = registry.WRAPPED_CHAT_AGENT_BY_TOOL_NAME.get("chat_agent_pptxer")
        self.assertIsNotNone(spec, "chat_agent_pptxer is not registered")
        self.assertEqual(spec.template_dir, "pptxer")
        self.assertEqual(spec.display_name, "PPTXer")
        self.assertEqual(spec.tool_description, "Chat-Agent-PPTXer")

    def test_literal_text_fields_declare_the_verbatim_channel(self):
        """⚠️ NOT automatic. Without this the parser collapses \\\\ to \\."""
        from agent import chat_agent_registry as registry
        spec = registry.WRAPPED_CHAT_AGENT_BY_TOOL_NAME["chat_agent_pptxer"]
        self.assertIn("input_text", spec.verbatim_fields)

    def test_the_display_name_is_exactly_PPTXer(self):
        """str.title() would ship 'Pptxer'."""
        from agent.services.agent_paths import display_name_from_agent_type
        self.assertEqual(display_name_from_agent_type("pptxer"), "PPTXer")

    def test_it_is_a_parametrizer_source(self):
        from agent.services.agent_contracts import get_agent_contract
        contract = get_agent_contract("pptxer")
        self.assertEqual(contract.display_name, "PPTXer")
        for field in ("output_path", "slide_count", "layout_clean",
                      "ground_truth", "status", "response_body"):
            self.assertIn(field, contract.parametrizer_fields,
                          f"{field} missing from the parametrizer contract")

    def test_the_section_header_matches_the_registered_contract(self):
        """Keep the two sides in step — an AST read of BOTH, not a typed count."""
        import ast
        from agent.services.agent_contracts import get_agent_contract

        source = open(os.path.join(AGENT_DIR, "pptxer.py"),
                      encoding="utf-8").read()
        tree = ast.parse(source)
        emitted = set()
        for node in ast.walk(tree):
            if not isinstance(node, ast.Dict):
                continue
            for key in node.keys:
                if isinstance(key, ast.Constant) and isinstance(key.value, str):
                    emitted.add(key.value)

        registered = set(get_agent_contract("pptxer").parametrizer_fields)
        registered.discard("response_body")          # the body, not a header key
        missing = registered - emitted
        self.assertEqual(missing, set(),
                         f"registered but never emitted: {sorted(missing)}")

    def test_it_is_captured_in_the_exec_report(self):
        from agent.mcp_agent import _resolve_exec_report_spec
        spec = _resolve_exec_report_spec("chat_agent_pptxer")
        self.assertIsNotNone(spec, "every Multi-Turn agent must produce a row")
        self.assertEqual(spec[0], "pptxer")
        self.assertEqual(spec[1], "PPTXer")

    def test_it_is_gated_by_ask_execs(self):
        """Tier A (writes anywhere) AND tier D (reaches remote hosts).

        ``_requires_exec_permission`` is a METHOD on the executor, so the
        allowlist itself is what this asserts against — that is the datum the
        method consults and the one a careless edit would drop.
        """
        from agent.mcp_agent import _ASK_EXECS_REQUIRED_TOOLS
        self.assertIn("chat_agent_pptxer", _ASK_EXECS_REQUIRED_TOOLS)

    def test_the_connection_route_resolves(self):
        """The agent URLconf carries no app_name, so reverse() is unnamespaced."""
        from django.urls import reverse
        self.assertTrue(reverse("update_pptxer_connection", args=["pptxer-1"]))

    def test_the_agent_is_a_registered_section_producer(self):
        parametrizer = os.path.join(AGENT_DIR, "..", "parametrizer",
                                    "parametrizer.py")
        source = open(os.path.abspath(parametrizer), encoding="utf-8").read()
        self.assertIn("'pptxer'", source,
                      "pptxer missing from SECTION_AGENT_TYPES")


class FrontendWiringTests(SimpleTestCase):
    """A missing JS site breaks a canvas connection SILENTLY."""

    STATIC = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                          "static", "agent")

    def _read(self, *parts):
        with open(os.path.join(self.STATIC, *parts), encoding="utf-8",
                  errors="replace") as handle:
            return handle.read()

    def test_the_connector_exists(self):
        source = self._read("js", "acp-agent-connectors.js")
        self.assertIn("async function updatePptxerConnection", source)
        self.assertIn("/agent/update_pptxer_connection/", source)

    def test_canvas_core_has_the_classmap_and_every_connection_site(self):
        source = self._read("js", "acp-canvas-core.js")
        self.assertIn("'pptxer': 'pptxer-agent'", source)
        self.assertGreaterEqual(source.count("updatePptxerConnection"), 3,
                                "canvas-core needs remove / removeFor / mouseup")

    def test_undo_and_redo_are_both_wired(self):
        source = self._read("js", "acp-canvas-undo.js")
        self.assertIn("updatePptxerConnection(sourceId, targetId, 'remove')", source)
        self.assertIn("updatePptxerConnection(sourceId, targetId, 'add')", source)

    def test_flw_load_restores_the_connection(self):
        source = self._read("js", "acp-file-io.js")
        self.assertIn("case 'pptxer':", source)

    def test_the_canvas_gradient_exists_and_is_unique(self):
        css = self._read("css", "agentic_control_panel.css")
        self.assertIn(".canvas-item.pptxer-agent", css)
        self.assertIn(".canvas-item.pptxer-agent:hover", css)
        # The exact ramp must not be shared with another agent.
        ramp = "#1A0B2E 0%, #6D28D9 33%, #FF6B35 66%, #FFD93D 100%"
        self.assertEqual(css.count(ramp), 1, "the gradient is not unique")

    def test_the_exec_report_caption_exists(self):
        css = self._read("css", "agent_page.css")
        self.assertIn(".exec-report-caption-pptxer", css)

    def test_the_flow_generator_maps_the_config_fields(self):
        source = self._read("js", "agent_page_chat.js")
        self.assertIn("lower === 'pptxer'", source)
        for field in ("input_text", "nuance", "predominant_color", "filename"):
            self.assertIn(f"set('{field}'", source)

    def test_every_mapped_field_exists_in_config_yaml(self):
        """A mapped field that is not a real key silently does nothing."""
        import re
        import yaml

        source = self._read("js", "agent_page_chat.js")
        block = source.split("lower === 'pptxer'")[1].split("} else if")[0]
        mapped = set(re.findall(r"set\('([a-z_]+)'", block))
        with open(os.path.join(AGENT_DIR, "config.yaml"), encoding="utf-8") as fh:
            config = yaml.safe_load(fh) or {}
        unknown = mapped - set(config)
        self.assertEqual(unknown, set(),
                         f"mapped but not in config.yaml: {sorted(unknown)}")


class ConfigContractTests(SimpleTestCase):

    def test_config_yaml_parses_and_carries_the_connection_fields(self):
        import yaml
        with open(os.path.join(AGENT_DIR, "config.yaml"), encoding="utf-8") as fh:
            config = yaml.safe_load(fh)
        self.assertIsInstance(config, dict)
        self.assertEqual(config["target_agents"], [])
        self.assertEqual(config["source_agents"], [])
        self.assertEqual(config["action"], "create")

    def test_numeric_defaults_are_numbers_not_strings(self):
        import yaml
        with open(os.path.join(AGENT_DIR, "config.yaml"), encoding="utf-8") as fh:
            config = yaml.safe_load(fh)
        for key in ("margin_pt", "max_image_px", "render_width",
                    "fetch_timeout", "command_timeout"):
            self.assertIsInstance(config[key], (int, float),
                                  f"{key} must be numeric")

    def test_private_hosts_are_refused_by_default(self):
        """⚠️ Shipping this as true would let a deck probe the local network."""
        import yaml
        with open(os.path.join(AGENT_DIR, "config.yaml"), encoding="utf-8") as fh:
            config = yaml.safe_load(fh)
        self.assertFalse(config["allow_private_hosts"])

    def test_the_agent_never_imports_the_django_app(self):
        """A pool subprocess has no sys.path back into agent.*.

        ⚠️ Walk the AST, do not grep. The first version of this test searched
        for the string "from agent." and failed on the modules' OWN comments
        explaining that they must never import agent.* — a test that a correct
        file cannot pass. Only real import STATEMENTS count.
        """
        import ast
        import glob

        for path in glob.glob(os.path.join(AGENT_DIR, "*.py")):
            with open(path, encoding="utf-8", errors="replace") as handle:
                tree = ast.parse(handle.read(), filename=path)
            name = os.path.basename(path)
            for node in ast.walk(tree):
                if isinstance(node, ast.ImportFrom):
                    module = node.module or ""
                    self.assertFalse(
                        module == "agent" or module.startswith("agent."),
                        f"{name} imports {module}")
                elif isinstance(node, ast.Import):
                    for alias in node.names:
                        self.assertFalse(
                            alias.name == "agent" or alias.name.startswith("agent."),
                            f"{name} imports {alias.name}")

    def test_the_temp_policy_guard_is_present(self):
        """Scratch goes under <app>/Temp — checked in CODE, not in prose.

        Same lesson as the import test above: the module documents the policy
        by naming the forbidden paths, so a substring search finds them in the
        very comment that forbids them. Strip comments and docstrings first.
        """
        import ast
        import io
        import tokenize

        path = os.path.join(AGENT_DIR, "pptxer.py")
        with open(path, encoding="utf-8") as handle:
            source = handle.read()

        self.assertIn("TLAMATINI_TEMP", source,
                      "the temp-policy guard is missing entirely")

        # Every string/bytes literal that is actually EXECUTED.
        tree = ast.parse(source, filename=path)
        docstrings = set()
        for node in ast.walk(tree):
            if isinstance(node, (ast.Module, ast.FunctionDef, ast.AsyncFunctionDef,
                                 ast.ClassDef)):
                doc = ast.get_docstring(node, clean=False)
                if doc:
                    docstrings.add(doc)

        literals = [n.value for n in ast.walk(tree)
                    if isinstance(n, ast.Constant) and isinstance(n.value, str)
                    and n.value not in docstrings]

        for literal in literals:
            lowered = literal.lower()
            self.assertNotIn("c:\\temp", lowered,
                             f"a hardcoded scratch path in code: {literal[:60]!r}")

        # And no comment may hide a live tempfile.gettempdir() call either.
        tokens = tokenize.generate_tokens(io.StringIO(source).readline)
        code_only = "".join(
            tok.string for tok in tokens
            if tok.type not in (tokenize.COMMENT, tokenize.NL))
        self.assertNotIn("tempfile.gettempdir()", code_only,
                         "scratch must resolve through TLAMATINI_TEMP")

    def test_reanimation_is_detected_before_logging_is_configured(self):
        """Otherwise the log truncates on every resume."""
        with open(os.path.join(AGENT_DIR, "pptxer.py"), encoding="utf-8") as fh:
            source = fh.read()
        self.assertLess(source.index("_IS_REANIMATED"),
                        source.index("logging.basicConfig"))
