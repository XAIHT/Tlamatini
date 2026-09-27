# ═══════════════════════════════════════════════════════════════════
#   ✦  T L A M A T I N I  ✦   —   "one who knows"
#
#   Created by  Angela López Mendoza   ·   @angelahack1
#   Developer · Architect · Creator of Tlamatini
#
#   Every line of this file was written by Angela López Mendoza.
# ═══════════════════════════════════════════════════════════════════
#   Tlamatini Author Banner — do not remove (releases scrub the name automatically)
"""The dossier deck: a 16:9 obsidian-and-jade presentation built from the same
content as the PDF.

Text is never left to PowerPoint's auto-fit. Every paragraph is measured with
the real TrueType font before it is placed (PowerPoint sets single-spaced lines
at exactly 1.2 x the font size), the largest size that fits its box is chosen,
and a box that cannot hold its text even at the minimum size stops the build.
Every intentional rectangle is recorded so the geometry audit can prove that
nothing leaves the slide and no text box overlaps another. The speaker notes
carry the complete narrative of each slide.
"""
from __future__ import annotations

import math
from dataclasses import dataclass, field
from functools import lru_cache

from lxml import etree
from PIL import ImageFont
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from dossier_theme import (FONTS, PALETTE, GlyphLedger, circular_portrait, deck_background, font_path, plain, runs,
                           split_fallback)

SW, SH = 13.333, 7.5
MX = 0.62
CONTENT_TOP = 2.28
CONTENT_BOTTOM = 6.78
CONTENT_W = SW - 2 * MX
LEDGER = GlyphLedger()
LINE = 1.2
WIDTH_SAFETY = 0.955
INSET_X, INSET_Y = 0.05, 0.03


class LayoutError(RuntimeError):
    pass


def rgb(name: str) -> RGBColor:
    return RGBColor.from_string(PALETTE.get(name, name).lstrip("#").upper())


# ── Measurement ─────────────────────────────────────────────────────
@lru_cache(maxsize=None)
def _pil_font(key: str, size10: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(font_path(key)), size=size10 * 4 / 10)


def measure(text: str, key: str, size: float, spacing: float = 0.0) -> float:
    """Width in points of text in a logical font, honouring symbol fallback."""
    total = 0.0
    for chunk, font in split_fallback(text, key):
        total += _pil_font(font, round(size * 10)).getlength(chunk) / 4
    return total + spacing * max(0, len(text) - 1)


@dataclass
class Run:
    text: str
    font: str
    size: float
    color: str
    spacing: float = 0.0


@dataclass
class Para:
    runs: list[Run]
    align: str = "left"
    space_before: float = 0.0
    bullet: bool = False


def _tokens(para: Para):
    for run in para.runs:
        pieces = run.text.replace("\n", " \n ").split(" ")
        for index, piece in enumerate(pieces):
            if piece == "\n":
                yield ("\n", run)
                continue
            word = piece + (" " if index < len(pieces) - 1 else "")
            if word:
                yield (word, run)


def para_lines(para: Para, max_w: float) -> int:
    lines = 1
    used = 0.0
    for word, run in _tokens(para):
        if word == "\n":
            lines += 1
            used = 0.0
            continue
        width = measure(word, run.font, run.size, run.spacing)
        trimmed = measure(word.rstrip(), run.font, run.size, run.spacing)
        if used and used + trimmed > max_w:
            lines += 1
            used = 0.0
        if trimmed > max_w:
            lines += math.ceil(trimmed / max_w) - 1
            used = trimmed % max_w
        else:
            used += width
    return lines


def block_height(paras: list[Para], max_w: float) -> float:
    total = 0.0
    for index, para in enumerate(paras):
        size = max(run.size for run in para.runs)
        total += para_lines(para, max_w) * size * LINE
        if index:
            total += para.space_before
    return total


def scaled(paras: list[Para], factor: float) -> list[Para]:
    return [Para([Run(r.text, r.font, r.size * factor, r.color, r.spacing * factor) for r in p.runs],
                 p.align, p.space_before * factor, p.bullet) for p in paras]


def marked(text: str, font: str, size: float, color: str, bold_font: str | None = None,
           code_color: str = "gold", bold_color: str | None = None) -> list[Run]:
    """Turn **bold** / `code` / _italic_ markup into runs."""
    out = []
    for fragment, kind in runs(text):
        if kind == "bold":
            out.append(Run(fragment, bold_font or "sans-semibold", size, bold_color or color))
        elif kind == "code":
            out.append(Run(fragment, "mono", size * 0.92, code_color))
        elif kind == "italic":
            out.append(Run(fragment, "sans-italic" if font.startswith("sans") else "serif-italic", size, color))
        else:
            out.append(Run(fragment, font, size, color))
    return out


# ── Slide canvas ────────────────────────────────────────────────────
@dataclass
class Rect:
    x: float
    y: float
    w: float
    h: float
    kind: str
    label: str
    parent: int | None = None


@dataclass
class SlideCanvas:
    slide: object
    number: int
    rects: list[Rect] = field(default_factory=list)

    def track(self, x, y, w, h, kind, label, parent=None) -> int:
        self.rects.append(Rect(x, y, w, h, kind, label, parent))
        return len(self.rects) - 1


def _alpha(fill_elm, alpha: float | None):
    if alpha is None:
        return
    clr = fill_elm.find(qn("a:srgbClr"))
    if clr is not None:
        node = etree.SubElement(clr, qn("a:alpha"))
        node.set("val", str(int(alpha * 100000)))


def shape_box(sc: SlideCanvas, x, y, w, h, fill=None, line=None, radius=0.08, line_w=0.75, alpha=None,
              kind="panel", label="panel", shape=MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE) -> int:
    shp = sc.slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.shadow.inherit = False
    if shape == MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = min(0.5, radius / max(0.01, min(w, h)))
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
        _alpha(shp.fill._xPr.find(qn("a:solidFill")), alpha)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(line_w)
    else:
        shp.line.fill.background()
    if shp.has_text_frame:
        shp.text_frame.text = ""
    return sc.track(x, y, w, h, kind, label)


def line(sc: SlideCanvas, x1, y1, x2, y2, color="edge", width=0.75, arrow_end=False):
    conn = sc.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = rgb(color)
    conn.line.width = Pt(width)
    if arrow_end:
        ln = conn.line._get_or_add_ln()
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle")
        tail.set("w", "med")
        tail.set("len", "med")
    sc.track(min(x1, x2), min(y1, y2), abs(x2 - x1), abs(y2 - y1), "line", "line")


def text(sc: SlideCanvas, x, y, w, h, paras: list[Para], anchor="top", min_scale=0.62, label="text",
         parent=None, wrap=True) -> float:
    """Place measured paragraphs; shrink proportionally until they fit, else fail."""
    avail_w = (w - 2 * INSET_X) * 72 * WIDTH_SAFETY
    avail_h = (h - 2 * INSET_Y) * 72
    chosen = None
    factor = 1.0
    while factor >= min_scale - 1e-9:
        candidate = scaled(paras, factor)
        need = block_height(candidate, avail_w if wrap else 10 ** 9)
        if not wrap:
            widest = max(sum(measure(r.text, r.font, r.size, r.spacing) for r in p.runs) for p in candidate)
            if widest > avail_w:
                factor -= 0.02
                continue
        if need <= avail_h - 0.5:
            chosen = candidate
            break
        factor -= 0.02
    if chosen is None:
        sample = plain("".join(r.text for r in paras[0].runs))[:60]
        raise LayoutError(f"Slide {sc.number}: text does not fit {w:.2f}x{h:.2f} in ({label}): {sample!r}")
    box = sc.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = Inches(INSET_X)
    tf.margin_top = tf.margin_bottom = Inches(INSET_Y)
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
    for index, para in enumerate(chosen):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[para.align]
        p.line_spacing = 1.0
        if index and para.space_before:
            p.space_before = Pt(para.space_before)
        for run in para.runs:
            for chunk, font in split_fallback(run.text, run.font):
                LEDGER.check(chunk, font)
                r = p.add_run()
                r.text = chunk
                family, bold, italic = FONTS[font][1], FONTS[font][2], FONTS[font][3]
                r.font.name = family
                r.font.size = Pt(round(run.size * 2) / 2)
                r.font.bold = bold
                r.font.italic = italic
                r.font.color.rgb = rgb(run.color)
                rpr = r._r.get_or_add_rPr()
                if run.spacing:
                    rpr.set("spc", str(int(run.spacing * 100)))
    sc.track(x, y, w, h, "text", label, parent)
    return factor


def one(text_value, font, size, color, align="left", spacing=0.0) -> list[Para]:
    return [Para([Run(text_value, font, size, color, spacing)], align)]


def picture(sc: SlideCanvas, path, x, y, w, h, kind="image", label="image"):
    sc.slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    sc.track(x, y, w, h, kind, label)


# ── Slide furniture ─────────────────────────────────────────────────
class Deck:
    def __init__(self, facts):
        self.facts = facts
        self.prs = Presentation()
        self.prs.slide_width = Inches(SW)
        self.prs.slide_height = Inches(SH)
        self.layout = self.prs.slide_layouts[6]
        self.canvases: list[SlideCanvas] = []
        self.backgrounds = {kind: deck_background(kind) for kind in ("cover", "divider", "content", "appendix")}

    def new(self, background="content", notes: str = "") -> SlideCanvas:
        slide = self.prs.slides.add_slide(self.layout)
        sc = SlideCanvas(slide, len(self.canvases) + 1)
        picture(sc, self.backgrounds[background], 0, 0, SW, SH, kind="bg", label="background")
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        self.canvases.append(sc)
        return sc

    def chrome(self, sc: SlideCanvas, label: str):
        line(sc, MX, 6.97, SW - MX, 6.97, "edge", 0.6)
        text(sc, MX, 7.02, 7.5, 0.3, one(f"TLAMATINI  ·  {label.upper()}", "sans-semibold", 8.5,
                                        "light_muted", spacing=1.2), label="footer")
        text(sc, SW - MX - 3.2, 7.02, 3.2, 0.3,
             one(f"v{self.facts['version']}   ·   {sc.number}", "sans-semibold", 8.5, "light_muted",
                 "right", 1.0), label="page")

    def header(self, sc: SlideCanvas, kicker: str, title: str, lede: str = "", accent: str = "jade"):
        shape_box(sc, MX, 0.5, 0.11, 0.11, fill=accent, radius=0.0, kind="deco", label="kicker mark",
                  shape=MSO_AUTO_SHAPE_TYPE.DIAMOND)
        text(sc, MX + 0.2, 0.40, 9.0, 0.32, one(kicker, "sans-semibold", 11, "copper", spacing=2.4),
             label="kicker")
        text(sc, MX - 0.03, 0.74, CONTENT_W, 0.72, one(title, "serif", 31, "ivory"), label="title", min_scale=0.7)
        shape_box(sc, MX, 1.47, 0.9, 0.04, fill=accent, radius=0.0, kind="deco", label="rule",
                  shape=MSO_AUTO_SHAPE_TYPE.RECTANGLE)
        if lede:
            text(sc, MX, 1.58, CONTENT_W - 0.4, 0.64,
                 [Para(marked(lede, "sans-semilight", 14.5, "light_muted", code_color="gold"))],
                 label="lede", min_scale=0.78)


# ── Components ──────────────────────────────────────────────────────
TONES = ["jade", "copper", "gold", "cyan", "magenta", "rose"]


def card(sc, x, y, w, h, title, body, tone="jade", title_font="sans-semibold", title_size=16, body_size=13,
         title_color="ivory"):
    parent = shape_box(sc, x, y, w, h, fill="panel", line="edge", radius=0.12, label="card")
    shape_box(sc, x + 0.18, y + 0.14, 0.55, 0.05, fill=tone, radius=0.0, kind="deco", label="accent",
              shape=MSO_AUTO_SHAPE_TYPE.RECTANGLE)
    inner_w = (w - 0.28 - 2 * INSET_X) * 72 * WIDTH_SAFETY
    lines = para_lines(Para(marked(title, title_font, title_size, title_color)), inner_w)
    title_h = min(h * 0.45, lines * title_size * LINE / 72 + 0.12)
    text(sc, x + 0.14, y + 0.24, w - 0.28, title_h, [Para(marked(title, title_font, title_size, title_color))],
         label="card title", parent=parent, min_scale=0.66)
    text(sc, x + 0.14, y + 0.24 + title_h, w - 0.28, h - title_h - 0.34,
         [Para(marked(body, "sans", body_size, "light_muted", code_color="gold"))],
         label="card body", parent=parent, min_scale=0.62)


def grid_dims(n: int) -> tuple[int, int]:
    if n <= 3:
        return n, 1
    if n == 4:
        return 2, 2
    if n <= 6:
        return 3, 2
    if n <= 8:
        return 4, 2
    if n <= 9:
        return 3, 3
    if n <= 12:
        return 4, 3
    return 4, math.ceil(n / 4)


def cards_grid(sc, points, x=MX, y=CONTENT_TOP, w=CONTENT_W, h=CONTENT_BOTTOM - CONTENT_TOP, gap=0.22,
               title_font="sans-semibold", title_size=16, body_size=13, tone=None):
    cols, rows = grid_dims(len(points))
    cw = (w - gap * (cols - 1)) / cols
    ch = (h - gap * (rows - 1)) / rows
    for index, (title, body) in enumerate(points):
        r, col = divmod(index, cols)
        card(sc, x + col * (cw + gap), y + r * (ch + gap), cw, ch, title, body, tone or TONES[index % 6],
             title_font, title_size, body_size)


def card_height(title, body, w, title_font, title_size, body_size) -> float:
    """The height a card needs so its title and body fit at their full sizes."""
    inner_w = (w - 0.28 - 2 * INSET_X) * 72 * WIDTH_SAFETY
    lines = para_lines(Para(marked(title, title_font, title_size, "ivory")), inner_w)
    title_h = lines * title_size * LINE / 72 + 0.12
    body_h = block_height([Para(marked(body, "sans", body_size, "light"))], inner_w) / 72 + 0.06 + 0.02
    return title_h + body_h + 0.34 + 0.06


def family_grid(sc, members, tone, title_size=18, body_size=14.5):
    """Agent cards sized to their content and centred in the content area."""
    cols, rows = grid_dims(len(members))
    gap = 0.24
    area_h = CONTENT_BOTTOM - CONTENT_TOP
    cw = (CONTENT_W - gap * (cols - 1)) / cols
    need = max(card_height(name, text_value, cw, "serif-bold", title_size, body_size)
               for name, text_value in members)
    ch = min((area_h - gap * (rows - 1)) / rows, max(need, 1.15))
    grid_h = rows * ch + gap * (rows - 1)
    top = CONTENT_TOP + (area_h - grid_h) / 2
    for index, (name, text_value) in enumerate(members):
        r, col = divmod(index, cols)
        card(sc, MX + col * (cw + gap), top + r * (ch + gap), cw, ch, name, text_value, tone, "serif-bold",
             title_size, body_size)


def list_rows(sc, points, x=MX, y=CONTENT_TOP, w=CONTENT_W, h=CONTENT_BOTTOM - CONTENT_TOP, head_w=3.0):
    n = len(points)
    gap = 0.14
    rh = (h - gap * (n - 1)) / n
    for index, (head, body) in enumerate(points):
        ry = y + index * (rh + gap)
        tone = TONES[index % 6]
        parent = shape_box(sc, x, ry, w, rh, fill="panel", line="edge", radius=0.1, label="row")
        shape_box(sc, x + 0.22, ry + rh / 2 - 0.07, 0.14, 0.14, fill=tone, radius=0.0, kind="deco",
                  label="bullet", shape=MSO_AUTO_SHAPE_TYPE.DIAMOND)
        text(sc, x + 0.5, ry + 0.06, head_w - 0.3, rh - 0.12, [Para(marked(head, "sans-semibold", 16, "ivory"))],
             anchor="middle", label="row head", parent=parent)
        text(sc, x + head_w + 0.3, ry + 0.06, w - head_w - 0.45, rh - 0.12,
             [Para(marked(body, "sans", 14, "light_muted", code_color="gold"))], anchor="middle",
             label="row body", parent=parent)


def table_block(sc, columns, rows, ratios, x=MX, y=CONTENT_TOP, w=CONTENT_W, h=CONTENT_BOTTOM - CONTENT_TOP,
                size=15.0, align=None):
    """A shape-drawn table whose rows are sized from measured text; raises if it cannot fit."""
    align = align or ["LEFT"] * len(columns)
    widths = [w * ratio for ratio in ratios]

    for attempt_size in [size - step * 0.5 for step in range(0, 11)]:
        heights = []
        head_h = 0.42
        for row in rows:
            need = 0.0
            for i, value in enumerate(row):
                avail = (widths[i] - 2 * INSET_X - 0.16) * 72 * WIDTH_SAFETY
                font = "sans-semibold" if i == 0 else "sans"
                paras = [Para(marked(str(value), font, attempt_size, "light"))]
                need = max(need, block_height(paras, avail) / 72)
            heights.append(max(0.30, need + 0.16))
        if head_h + sum(heights) <= h:
            size = attempt_size
            break
    else:
        raise LayoutError(f"Slide {sc.number}: table rows do not fit; split the table")
    parent = shape_box(sc, x, y, w, head_h + sum(heights), fill="panel", line="edge", radius=0.08, label="table")
    shape_box(sc, x, y, w, head_h, fill="stone", radius=0.08, kind="deco", label="table head")
    cx = x
    for i, name in enumerate(columns):
        text(sc, cx + 0.08, y + 0.04, widths[i] - 0.16, head_h - 0.08,
             [Para([Run(name.upper(), "sans-semibold", 10, "gold", 1.2)],
                   "right" if align[i] == "RIGHT" else "left")], anchor="middle", label="th", parent=parent)
        cx += widths[i]
    ry = y + head_h
    for r_index, (row, rh) in enumerate(zip(rows, heights)):
        total = str(row[0]).strip().lower() == "total"
        if r_index % 2 == 1 or total:
            shape_box(sc, x + 0.02, ry, w - 0.04, rh, fill="jade" if total else "panel2", radius=0.0,
                      kind="deco", label="stripe", alpha=0.22 if total else None,
                      shape=MSO_AUTO_SHAPE_TYPE.RECTANGLE)
        cx = x
        for i, value in enumerate(row):
            font = "sans-semibold" if (i == 0 or total) else "sans"
            color = "ivory" if (i == 0 or total) else "light_muted"
            text(sc, cx + 0.08, ry + 0.04, widths[i] - 0.16, rh - 0.08,
                 [Para(marked(str(value), font, size, color, code_color="gold"),
                       "right" if align[i] == "RIGHT" else "left")], anchor="middle", label="td",
                 parent=parent, min_scale=0.95)
            cx += widths[i]
        ry += rh
    return head_h + sum(heights)


def table_rows_fit(columns, rows, ratios, w, h, size) -> int:
    """How many leading rows fit in height h (used to paginate long tables)."""
    widths = [w * ratio for ratio in ratios]
    used = 0.42
    for count, row in enumerate(rows):
        need = 0.0
        for i, value in enumerate(row):
            avail = (widths[i] - 2 * INSET_X - 0.16) * 72 * WIDTH_SAFETY
            font = "sans-semibold" if i == 0 else "sans"
            need = max(need, block_height([Para(marked(str(value), font, size, "light"))], avail) / 72)
        used += max(0.30, need + 0.16)
        if used > h:
            return count
    return len(rows)


# ── Visuals ─────────────────────────────────────────────────────────
def visual_planes(sc, x=MX, y=CONTENT_TOP, w=CONTENT_W, h=CONTENT_BOTTOM - CONTENT_TOP):
    spec = [("jade", "Local control plane", "YOUR MACHINE",
             ["Django, Channels and Daphne", "Chat, canvases and dialogs", "SQLite database and settings",
              "Agent programs and permissions"]),
            ("gold", "Local retrieval", "YOUR MACHINE",
             ["nomic-embed-text embeddings", "FAISS vectors and BM25", "Context budgeting",
              "Binary-content guard"]),
            ("copper", "Cloud reasoning", "OLLAMA CLOUD",
             ["Chat and tool calling", "Long context", "Vision observers and mergers", "Multi-Turn planning"])]
    gap = 0.55
    cw = (w - 2 * gap) / 3
    for index, (tone, title, place, items) in enumerate(spec):
        cx = x + index * (cw + gap)
        parent = shape_box(sc, cx, y, cw, h - 0.45, fill="panel", line=tone, radius=0.14, line_w=1.25,
                           label="plane")
        shape_box(sc, cx, y, cw, 0.62, fill=tone, radius=0.14, kind="deco", label="plane head")
        text(sc, cx + 0.2, y + 0.08, cw - 0.4, 0.48, one(title, "sans-semibold", 17, "obsidian"),
             anchor="middle", label="plane title", parent=parent)
        paras = [Para([Run("▪  " + item, "sans", 14.5, "light")], space_before=9) for item in items]
        text(sc, cx + 0.25, y + 0.85, cw - 0.45, h - 1.5, paras, label="plane items", parent=parent)
        text(sc, cx, y + h - 0.38, cw, 0.34, one(place, "sans-semibold", 10.5, tone, "center", 2.0),
             label="plane place")
        if index < 2:
            ay = y + (h - 0.45) / 2
            line(sc, cx + cw + 0.08, ay - 0.1, cx + cw + gap - 0.08, ay - 0.1, "copper", 1.5, True)
            line(sc, cx + cw + gap - 0.08, ay + 0.12, cx + cw + 0.08, ay + 0.12, "jade", 1.5, True)


def visual_layers(sc, layers, x=MX, y=CONTENT_TOP, w=CONTENT_W, h=CONTENT_BOTTOM - CONTENT_TOP):
    band = 0.4
    gap = 0.1
    inner_h = h - 2 * band - 2 * gap
    row_h = (inner_h - gap * (len(layers) - 1)) / len(layers)
    shape_box(sc, x, y, w, band, fill="stone", radius=0.08, label="browser band")
    text(sc, x + 0.2, y, w - 0.4, band, [Para([Run("BROWSER   ", "sans-semibold", 11, "gold", 1.6),
                                             Run("Chat page  ·  Agentic Control Panel  ·  Prompt Flow Panel"
                                                 "  ·  WebSockets", "sans", 13, "light")])],
         anchor="middle", label="browser")
    ry = y + band + gap
    for index, (title, desc) in enumerate(layers):
        tone = TONES[index % 6]
        parent = shape_box(sc, x, ry, w, row_h, fill="panel", line="edge", radius=0.08, label="layer")
        shape_box(sc, x, ry, 0.62, row_h, fill=tone, radius=0.08, kind="deco", label="layer tab")
        text(sc, x, ry, 0.62, row_h, one(str(index + 1), "serif-bold", 22, "obsidian", "center"),
             anchor="middle", label="layer number", parent=parent)
        text(sc, x + 0.8, ry + 0.04, 3.1, row_h - 0.08, one(title, "sans-semibold", 16.5, "ivory"),
             anchor="middle", label="layer title", parent=parent)
        text(sc, x + 4.0, ry + 0.04, w - 4.2, row_h - 0.08,
             [Para(marked(desc, "sans", 14, "light_muted", code_color="gold"))], anchor="middle",
             label="layer text", parent=parent)
        ry += row_h + gap
    shape_box(sc, x, ry, w, band, fill="stone", radius=0.08, label="models band")
    text(sc, x + 0.2, ry, w - 0.4, band, [Para([Run("MODELS   ", "sans-semibold", 11, "gold", 1.6),
                                              Run("Ollama (local and cloud)  ·  Anthropic Claude  ·  Qwen "
                                                  "vision", "sans", 13, "light")])], anchor="middle", label="models")


def visual_steps(sc, steps, x=MX, y=CONTENT_TOP, w=CONTENT_W, h=CONTENT_BOTTOM - CONTENT_TOP, cols=None):
    n = len(steps)
    cols = cols or (4 if n > 6 else (n if n <= 5 else 3))
    rows = math.ceil(n / cols)
    gap_x, gap_y = 0.34, 0.18
    cw = (w - gap_x * (cols - 1)) / cols
    ch = (h - gap_y * (rows - 1)) / rows
    for index, (title, desc) in enumerate(steps):
        r, col = divmod(index, cols)
        cx, cy = x + col * (cw + gap_x), y + r * (ch + gap_y)
        parent = shape_box(sc, cx, cy, cw, ch, fill="panel", line="edge", radius=0.12, label="step")
        shape_box(sc, cx + 0.14, cy + 0.12, 0.42, 0.42, fill="stone", line="gold", radius=0.21, line_w=1.0,
                  kind="deco", label="step badge", shape=MSO_AUTO_SHAPE_TYPE.OVAL)
        text(sc, cx + 0.14, cy + 0.12, 0.42, 0.42, one(str(index + 1), "serif-bold", 14, "gold", "center"),
             anchor="middle", label="step number", parent=parent)
        text(sc, cx + 0.64, cy + 0.08, cw - 0.76, 0.5, one(title, "sans-semibold", 16, "ivory"), anchor="middle",
             label="step title", parent=parent, min_scale=0.7)
        text(sc, cx + 0.14, cy + 0.62, cw - 0.26, ch - 0.7,
             [Para(marked(desc, "sans", 13, "light_muted", code_color="gold"))], label="step text",
             parent=parent)
        if col < cols - 1 and index + 1 < n:
            line(sc, cx + cw + 0.04, cy + 0.39, cx + cw + gap_x - 0.04, cy + 0.39, "copper", 1.4, True)


def visual_ladder(sc, rungs, x=MX, y=CONTENT_TOP, w=CONTENT_W, h=CONTENT_BOTTOM - CONTENT_TOP, cols=2):
    per_col = math.ceil(len(rungs) / cols)
    gap = 0.4
    cw = (w - gap * (cols - 1)) / cols
    rh_gap = 0.14
    rh = (h - rh_gap * (per_col - 1)) / per_col
    for index, (num, name, desc) in enumerate(rungs):
        col, r = divmod(index, per_col)
        cx, cy = x + col * (cw + gap), y + r * (rh + rh_gap)
        parent = shape_box(sc, cx + 0.34, cy, cw - 0.34, rh, fill="panel", line="edge", radius=0.1, label="rung")
        shape_box(sc, cx, cy + rh / 2 - 0.26, 0.52, 0.52, fill="stone", line="jade", radius=0.26, line_w=1.2,
                  kind="deco", label="rung badge", shape=MSO_AUTO_SHAPE_TYPE.OVAL)
        text(sc, cx, cy + rh / 2 - 0.26, 0.52, 0.52, one(num, "serif-bold", 15, "gold", "center"),
             anchor="middle", label="rung number")
        text(sc, cx + 0.62, cy + 0.04, 1.95, rh - 0.08, one(name, "mono-bold", 14.5, "jade"), anchor="middle",
             label="rung name", parent=parent, min_scale=0.7)
        text(sc, cx + 2.6, cy + 0.04, cw - 2.72, rh - 0.08,
             [Para(marked(desc, "sans", 13, "light_muted", code_color="gold"))], anchor="middle",
             label="rung text", parent=parent)


def visual_cascade(sc, stages, x=MX, y=CONTENT_TOP, w=CONTENT_W, h=CONTENT_BOTTOM - CONTENT_TOP):
    per_row = 4
    gap = 0.34
    bw = (w - gap * (per_row - 1)) / per_row
    bh = 1.25
    for index, (name, note) in enumerate(stages):
        r, col = divmod(index, per_row)
        bx, by = x + col * (bw + gap), y + 0.2 + r * (bh + 0.55)
        tone = "jade" if name in ("BOM", "Empty") else ("copper" if index in (0, 4, 5, 6, 7) else "gold")
        parent = shape_box(sc, bx, by, bw, bh, fill="panel", line=tone, radius=0.12, line_w=1.4, label="stage")
        text(sc, bx + 0.18, by + 0.14, bw - 0.36, 0.5, [Para([Run(f"{index + 1}   ", "serif-bold", 16, tone),
                                                             Run(name, "sans-semibold", 16, "ivory")])],
             label="stage name", parent=parent)
        text(sc, bx + 0.18, by + 0.66, bw - 0.36, bh - 0.78, one(note, "sans", 13, "light_muted"),
             label="stage note", parent=parent)
        if col < per_row - 1:
            line(sc, bx + bw + 0.04, by + bh / 2, bx + bw + gap - 0.04, by + bh / 2, "copper", 1.4, True)
    text(sc, x, y + h - 0.55, w, 0.5, [Para(marked(
        "Copper stages can drop a binary file; jade stages can only keep a file as text. Any doubt keeps the "
        "file, and every omission is logged under `--- [BINARY-GUARD]`.", "sans-italic", 13, "light_muted",
        code_color="gold"))], label="cascade note")


def _op_shape(sc, kind, x, y, w, h, tone):
    shapes = {"prompt": MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, "programmed": MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
              "decision": MSO_AUTO_SHAPE_TYPE.DIAMOND, "feed": MSO_AUTO_SHAPE_TYPE.ISOSCELES_TRIANGLE,
              "flush": MSO_AUTO_SHAPE_TYPE.FLOWCHART_MERGE, "clean": MSO_AUTO_SHAPE_TYPE.TRAPEZOID,
              "commentary": MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGULAR_CALLOUT}
    shape_box(sc, x, y, w, h, fill=tone, line=tone, radius=0.08, alpha=0.25, kind="deco", label="op glyph",
              shape=shapes[kind])
    if kind == "programmed":
        shape_box(sc, x + w - 0.3, y + 0.06, 0.24, 0.24, fill="panel", line=tone, radius=0.12, kind="deco",
                  label="clock", shape=MSO_AUTO_SHAPE_TYPE.OVAL)


def visual_operations(sc, ops, x=MX, y=CONTENT_TOP, w=CONTENT_W, h=CONTENT_BOTTOM - CONTENT_TOP):
    tones = {"prompt": "jade", "programmed": "jade", "decision": "gold", "feed": "cyan", "flush": "cyan",
             "clean": "copper", "commentary": "magenta"}
    cols = 4
    rows = math.ceil(len(ops) / cols)
    gap = 0.24
    cw = (w - gap * (cols - 1)) / cols
    ch = (h - gap * (rows - 1)) / rows
    for index, (kind, name, desc) in enumerate(ops):
        r, col = divmod(index, cols)
        cx, cy = x + col * (cw + gap), y + r * (ch + gap)
        parent = shape_box(sc, cx, cy, cw, ch, fill="panel", line="edge", radius=0.12, label="op card")
        _op_shape(sc, kind, cx + 0.2, cy + 0.2, 0.95, 0.6, tones[kind])
        text(sc, cx + 1.3, cy + 0.16, cw - 1.42, 0.68, one(name, "sans-semibold", 16, "ivory"), anchor="middle",
             label="op name", parent=parent, min_scale=0.7)
        text(sc, cx + 0.2, cy + 0.95, cw - 0.4, ch - 1.07, one(desc, "sans", 13, "light_muted"), label="op text",
             parent=parent)
    if len(ops) % cols:
        col = len(ops) % cols
        r = rows - 1
        cx, cy = x + col * (cw + gap), y + r * (ch + gap)
        parent = shape_box(sc, cx, cy, w - col * (cw + gap), ch, fill="stone", line="gold", radius=0.12,
                           label="format card")
        text(sc, cx + 0.25, cy + 0.2, w - col * (cw + gap) - 0.5, ch - 0.4, [
            Para([Run(".fpmt", "mono-bold", 22, "gold")]),
            Para(marked("A portable JSON diagram: `tlamatini-prompting-flow`, version 1. Opening a file never "
                        "runs it; Validate, then Play.", "sans", 13, "light", code_color="gold"), space_before=8)],
             label="format text", parent=parent)


def visual_metrics(sc, tiles, x=MX, y=CONTENT_TOP + 0.05, w=CONTENT_W, h=CONTENT_BOTTOM - CONTENT_TOP - 0.05):
    cols, rows = 4, math.ceil(len(tiles) / 4)
    gap = 0.26
    tw = (w - gap * (cols - 1)) / cols
    th = (h - gap * (rows - 1)) / rows
    for index, (value, label) in enumerate(tiles):
        r, col = divmod(index, cols)
        tx, ty = x + col * (tw + gap), y + r * (th + gap)
        tone = TONES[index % 6]
        parent = shape_box(sc, tx, ty, tw, th, fill="panel", line="edge", radius=0.14, label="tile")
        shape_box(sc, tx + 0.28, ty + 0.3, 0.7, 0.05, fill=tone, radius=0.0, kind="deco", label="tile accent",
                  shape=MSO_AUTO_SHAPE_TYPE.RECTANGLE)
        text(sc, tx + 0.22, ty + 0.42, tw - 0.44, th - 1.06, one(value, "serif-bold", 46, "ivory"), anchor="middle",
             label="tile value", parent=parent, min_scale=0.5)
        text(sc, tx + 0.24, ty + th - 0.6, tw - 0.48, 0.44, one(label.upper(), "sans-semibold", 11.5, tone,
                                                                 spacing=1.6), label="tile label", parent=parent)


def visual_family_tiles(sc, x=MX, y=CONTENT_TOP, w=CONTENT_W, h=CONTENT_BOTTOM - CONTENT_TOP):
    from dossier_content import AGENT_FAMILIES

    cols, rows = 3, math.ceil(len(AGENT_FAMILIES) / 3)
    gap = 0.2
    tw = (w - gap * (cols - 1)) / cols
    th = (h - gap * (rows - 1)) / rows
    for index, (name, tone, blurb, members) in enumerate(AGENT_FAMILIES):
        r, col = divmod(index, cols)
        tx, ty = x + col * (tw + gap), y + r * (th + gap)
        parent = shape_box(sc, tx, ty, tw, th, fill="panel", line="edge", radius=0.1, label="family tile")
        shape_box(sc, tx, ty + 0.12, 0.06, th - 0.24, fill=tone, radius=0.0, kind="deco", label="family bar",
                  shape=MSO_AUTO_SHAPE_TYPE.RECTANGLE)
        text(sc, tx + 0.16, ty + 0.05, 0.86, th - 0.1, one(str(len(members)), "serif-bold", 30, "ivory", "center"),
             anchor="middle", label="family count", parent=parent)
        text(sc, tx + 1.05, ty + 0.08, tw - 1.15, th * 0.5 - 0.08, one(name, "sans-semibold", 14.5, tone),
             anchor="bottom", label="family name", parent=parent, min_scale=0.7)
        text(sc, tx + 1.05, ty + th * 0.5, tw - 1.15, th * 0.5 - 0.08, one(blurb, "sans", 11.5, "light_muted"),
             label="family blurb", parent=parent)


def visual_chain(sc, items, caption, x=MX, y=CONTENT_TOP, w=CONTENT_W, h=1.0):
    n = len(items)
    gap = 0.36
    bw = (w - gap * (n - 1)) / n
    for index, name in enumerate(items):
        bx = x + index * (bw + gap)
        tone = "jade" if index in (0, n - 1) else "stone"
        parent = shape_box(sc, bx, y, bw, 0.68, fill=tone, line="jade" if tone == "stone" else None, radius=0.34,
                           label="chain pill")
        text(sc, bx + 0.08, y + 0.06, bw - 0.16, 0.56, one(name, "sans-semibold", 15,
                                                           "obsidian" if tone == "jade" else "gold", "center"),
             anchor="middle", label="chain name", parent=parent, min_scale=0.6)
        if index < n - 1:
            line(sc, bx + bw + 0.04, y + 0.34, bx + bw + gap - 0.04, y + 0.34, "copper", 1.5, True)
    if caption:
        text(sc, x, y + 0.78, w, 0.36, one(caption, "sans-italic", 12.5, "light_muted"), label="chain caption")


def visual_timeline(sc, rows, x=MX, y=CONTENT_TOP, w=CONTENT_W, h=CONTENT_BOTTOM - CONTENT_TOP):
    n = len(rows)
    per_col = math.ceil(n / 2)
    gap = 0.5
    cw = (w - gap) / 2
    rh = h / per_col
    for index, row in enumerate(rows):
        col, r = divmod(index, per_col)
        cx, cy = x + col * (cw + gap), y + r * rh
        last = index == n - 1
        if r < per_col - 1 and index < n - 1:
            line(sc, cx + 1.52, cy + rh / 2 + 0.1, cx + 1.52, cy + rh * 1.5 - 0.1, "copper", 1.2)
        shape_box(sc, cx + 1.44, cy + rh / 2 - 0.08, 0.16, 0.16, fill="jade" if last else "copper", radius=0.0,
                  kind="deco", label="tick", shape=MSO_AUTO_SHAPE_TYPE.DIAMOND)
        text(sc, cx, cy + 0.02, 1.36, rh - 0.04, one(row["date"], "mono", 12, "light_muted"),
             anchor="middle", label="tl date")
        text(sc, cx + 1.72, cy + 0.02, 1.25, rh - 0.04, one(row["tag"], "serif-bold", 16,
                                                            "gold" if last else "ivory"),
             anchor="middle", label="tl tag")
        text(sc, cx + 3.0, cy + 0.02, cw - 3.0, rh - 0.04, one(row["subject"], "sans", 12, "light_muted"),
             anchor="middle", label="tl subject", min_scale=0.6)


def visual_bars(sc, rows, x=MX, y=CONTENT_TOP, w=CONTENT_W, h=CONTENT_BOTTOM - CONTENT_TOP, title=""):
    data = CategoryChartData()
    data.categories = [name for name, _ in rows][::-1]
    data.add_series(title or "Value", [value for _, value in rows][::-1])
    frame = sc.slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), data)
    sc.track(x, y, w, h, "chart", "chart")
    chart = frame.chart
    chart.has_legend = False
    chart.has_title = False
    chart.font.name = "Segoe UI"
    chart.font.size = Pt(12)
    chart.font.color.rgb = rgb("light")
    plot = chart.plots[0]
    plot.gap_width = 45
    plot.vary_by_categories = False
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format = "#,##0"
    labels.number_format_is_linked = False
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    labels.font.size = Pt(11)
    labels.font.color.rgb = rgb("gold")
    series = plot.series[0]
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = rgb("jade")
    value_axis = chart.value_axis
    value_axis.visible = False
    value_axis.has_major_gridlines = False
    category_axis = chart.category_axis
    category_axis.format.line.color.rgb = rgb("edge")
    category_axis.tick_labels.font.size = Pt(12)
    category_axis.tick_labels.font.color.rgb = rgb("light")


# ── Slide builders ──────────────────────────────────────────────────
def cover_slide(deck: Deck):
    f = deck.facts
    sc = deck.new("cover", notes="Tlamatini: the complete project dossier. Created by Angela López Mendoza.")
    portrait = circular_portrait()
    d = 4.25
    picture(sc, portrait, SW * 0.30 - d / 2, SH * 0.52 - d / 2, d, d, label="portrait")
    x = 7.05
    text(sc, x, 0.98, 5.8, 0.36, one("THE COMPLETE PROJECT DOSSIER", "sans-semibold", 12, "copper", spacing=3.4),
         label="cover kicker")
    text(sc, x - 0.06, 1.38, 5.9, 1.33, one("Tlamatini", "serif", 76, "ivory"), label="cover title")
    text(sc, x, 2.74, 5.8, 0.55, one("the one who knows", "serif-italic", 26, "jade"), label="cover tagline")
    text(sc, x, 3.42, 5.7, 1.1, [Para(marked(
        "A self-hosted AI developer assistant with a visual agent-workflow designer, a universal MCP client, "
        "and the reach to drive real boards, engines and networks.", "sans-light", 16.5, "light"))],
         label="cover subtitle")
    metrics = [(str(f["agents"]), "agent types"), (str(f["tools"]), "Multi-Turn tools"),
               (str(f["skills"]), "skills"), (f"v{f['version']}", "release")]
    mw = 5.8 / 4
    for index, (value, label) in enumerate(metrics):
        mx = x + index * mw
        text(sc, mx, 4.78, mw - 0.1, 0.6, one(value, "serif-bold", 30, "gold"), label="cover metric")
        text(sc, mx, 5.40, mw - 0.1, 0.3, one(label.upper(), "sans-semibold", 9.5, "light_muted", spacing=1.2),
             label="cover metric label")
        if index:
            line(sc, mx - 0.1, 4.86, mx - 0.1, 5.6, "edge", 0.9)
    text(sc, x, 6.12, 5.8, 0.34, one("Created by Angela López Mendoza  ·  @angelahack1", "sans-semibold",
                                     13, "ivory"), label="cover credit")
    text(sc, x, 6.46, 5.8, 0.3, one(f"{f['generated_date']}  ·  github.com/XAIHT/Tlamatini  ·  MIT",
                                    "sans", 10.5, "light_muted"), label="cover date")


def agenda_slide(deck: Deck, chapters):
    sc = deck.new("content", notes="What this dossier covers, chapter by chapter.")
    deck.header(sc, "CONTENTS", "What this dossier covers",
                "Nine chapters from identity to inventory, and a complete file tree at the end.")
    items = [(ch.numeral, ch.title, ch.tagline, ch.accent) for ch in chapters]
    items.append(("A", "The complete tracked file tree", "Every file Git tracks, drawn as a tree.", "gold"))
    cols = 2
    per_col = math.ceil(len(items) / cols)
    gap = 0.4
    cw = (CONTENT_W - gap) / cols
    rh = (CONTENT_BOTTOM - CONTENT_TOP) / per_col
    for index, (numeral, title, tagline, accent) in enumerate(items):
        col, r = divmod(index, per_col)
        cx, cy = MX + col * (cw + gap), CONTENT_TOP + r * rh
        text(sc, cx, cy, 0.8, rh - 0.06, one(numeral, "serif", 26, accent, "center"), anchor="middle",
             label="agenda numeral")
        text(sc, cx + 0.9, cy + 0.02, cw - 0.95, 0.40, one(title, "sans-semibold", 16, "ivory"), label="agenda title",
             min_scale=0.7)
        text(sc, cx + 0.9, cy + 0.42, cw - 0.95, rh - 0.46, one(tagline, "sans", 12, "light_muted"),
             label="agenda tagline", min_scale=0.62)
    deck.chrome(sc, "Contents")


def divider_slide(deck: Deck, chapter, label_override: str | None = None):
    sc = deck.new("divider", notes=f"Chapter {chapter.numeral}: {chapter.title}. {chapter.tagline}")
    text(sc, 0.9, 0.95, 5.0, 0.4, one("CHAPTER" if chapter.numeral[0] in "IVX" else "APPENDIX", "sans-semibold",
                                      13, chapter.accent, spacing=3.4), label="div kicker")
    text(sc, 0.82, 1.38, 5.5, 1.85, one(chapter.numeral, "serif", 120, "copper"), label="div numeral", min_scale=0.5)
    text(sc, 0.9, 3.25, 7.4, 1.4, one(label_override or chapter.title, "serif", 46, "ivory"), label="div title",
         min_scale=0.6)
    shape_box(sc, 0.92, 4.72, 1.1, 0.05, fill=chapter.accent, radius=0.0, kind="deco", label="div rule",
              shape=MSO_AUTO_SHAPE_TYPE.RECTANGLE)
    text(sc, 0.9, 4.86, 7.0, 0.8, [Para(marked(chapter.tagline, "sans-light", 18, "light"))], label="div tagline",
         min_scale=0.7)
    items = [s.title for s in chapter.sections if s.deck != "none"]
    if 0 < len(items) <= 6:
        for index, name in enumerate(items):
            col, row = divmod(index, 3)
            x, y = 0.9 + col * 3.7, 5.74 + row * 0.3
            shape_box(sc, x, y + 0.1, 0.1, 0.1, fill=chapter.accent, radius=0.0, kind="deco", label="div bullet",
                      shape=MSO_AUTO_SHAPE_TYPE.DIAMOND)
            text(sc, x + 0.2, y, 3.4, 0.3, one(name, "sans", 12, "light_muted"), label="div section",
                 min_scale=0.7)
    return sc


def section_slides(deck: Deck, chapter, section):
    f = deck.facts
    notes = "\n\n".join([plain(section.lede)] + [plain(t) for t in section.body]
                        + [f"{h}: {plain(b)}" for h, b in section.points])
    label = f"{chapter.numeral} · {chapter.title}"
    points = section.deck_points or section.points
    layout = section.deck

    def start(kicker=None, title=None, lede=None):
        sc = deck.new("content", notes=notes)
        deck.header(sc, kicker or section.kicker, title or section.title,
                    section.lede if lede is None else lede, chapter.accent)
        return sc

    if layout == "identity":
        sc = deck.new("content", notes=notes)
        portrait = circular_portrait()
        picture(sc, portrait, 0.7, 1.05, 4.2, 4.2, label="identity portrait")
        text(sc, 0.7, 5.35, 4.2, 0.5, one("“one who knows”", "serif-italic", 20, "jade", "center"),
             label="identity quote")
        text(sc, 0.7, 5.88, 4.2, 0.36, one("NAHUATL", "sans-semibold", 11, "copper", "center", 3.0),
             label="identity lang")
        x = 5.45
        text(sc, x, 0.48, 7.2, 0.3, one(section.kicker, "sans-semibold", 11, "copper", spacing=2.4),
             label="kicker")
        text(sc, x - 0.03, 0.78, 7.3, 0.8, one(section.title, "serif", 31, "ivory"), label="title")
        text(sc, x, 1.62, 7.2, 1.38, [Para(marked(section.lede, "sans-semilight", 15, "light"))], label="identity lede")
        for index, (head, body) in enumerate(section.points):
            ry = 3.18 + index * 0.84
            parent = shape_box(sc, x, ry, 7.2, 0.7, fill="panel", line="edge", radius=0.1, label="fact row")
            text(sc, x + 0.22, ry + 0.06, 1.7, 0.58, one(head.upper(), "sans-semibold", 11, TONES[index],
                                                          spacing=1.6), anchor="middle", label="fact head",
                 parent=parent)
            text(sc, x + 2.0, ry + 0.06, 5.05, 0.58, [Para(marked(body, "sans", 15, "ivory"))], anchor="middle",
                 label="fact body", parent=parent)
        deck.chrome(sc, label)
        return
    if layout == "metrics":
        sc = start()
        from dossier_pdf import metric_tiles

        visual_metrics(sc, metric_tiles(f))
        deck.chrome(sc, label)
        return
    if layout == "family":
        tone = section.visual_data.get("accent", "jade")
        members = section.points
        chunks = [members] if len(members) <= 12 else [members[:7], members[7:]]
        for part, chunk in enumerate(chunks):
            suffix = f"  ({part + 1}/{len(chunks)})" if len(chunks) > 1 else ""
            sc = start(f"AGENT FAMILY  ·  {len(members)} AGENTS", section.title + suffix, section.lede)
            sizes = (21, 16.5) if len(chunk) <= 6 else ((19, 15) if len(chunk) <= 8 else (17.5, 14))
            family_grid(sc, chunk, tone, *sizes)
            deck.chrome(sc, label)
        return
    if layout == "table" and section.table:
        t = section.table
        rows = list(t["rows"])
        pages = []
        while rows:
            count = max(1, table_rows_fit(t["columns"], rows, t["widths"], CONTENT_W,
                                          CONTENT_BOTTOM - CONTENT_TOP, 10.5))
            pages.append(rows[:count])
            rows = rows[count:]
        for part, chunk in enumerate(pages):
            suffix = f"  ({part + 1}/{len(pages)})" if len(pages) > 1 else ""
            sc = start(title=section.title + suffix)
            table_block(sc, t["columns"], chunk, t["widths"], align=t.get("align"))
            deck.chrome(sc, label)
        if section.callout:
            callout_slide(deck, chapter, section, notes, label)
        return
    if layout == "visual":
        sc = start()
        kind = section.visual
        data = section.visual_data
        if kind == "planes":
            visual_planes(sc)
        elif kind == "layers":
            visual_layers(sc, data["layers"])
        elif kind == "steps":
            if section.code:
                visual_steps(sc, data["steps"], h=2.05, cols=5)
                code_panel(sc, section.code, MX, CONTENT_TOP + 2.3, CONTENT_W, CONTENT_BOTTOM - CONTENT_TOP - 2.3)
            else:
                visual_steps(sc, data["steps"])
        elif kind == "ladder":
            visual_ladder(sc, data["rungs"])
        elif kind == "cascade":
            visual_cascade(sc, data["stages"])
        elif kind == "operations":
            visual_operations(sc, data["ops"])
        elif kind == "timeline":
            visual_timeline(sc, f["timeline"])
        elif kind == "language_bars":
            from dossier_pdf import language_series

            visual_bars(sc, language_series(f, 14), title="Effective lines")
        elif kind == "family_wheel":
            visual_family_tiles(sc)
        deck.chrome(sc, label)
        if section.callout:
            callout_slide(deck, chapter, section, notes, label)
        return
    if layout == "split":
        sc = start()
        left_w = 5.7
        list_rows(sc, points, w=left_w, head_w=2.0)
        rx = MX + left_w + 0.4
        rw = CONTENT_W - left_w - 0.4
        right_panel(sc, section, rx, CONTENT_TOP, rw, CONTENT_BOTTOM - CONTENT_TOP)
        deck.chrome(sc, label)
        if section.callout:
            callout_slide(deck, chapter, section, notes, label)
        return
    if layout == "list":
        sc = start()
        list_rows(sc, points, head_w=3.2)
        deck.chrome(sc, label)
        if section.callout:
            callout_slide(deck, chapter, section, notes, label)
        return
    # cards (default)
    sc = start()
    cards_grid(sc, points)
    deck.chrome(sc, label)
    if section.callout:
        callout_slide(deck, chapter, section, notes, label)


def code_panel(sc, code, x, y, w, h):
    parent = shape_box(sc, x, y, w, h, fill="obsidian", line="edge", radius=0.12, label="code panel")
    for index, tone in enumerate(("rose", "gold", "jade")):
        shape_box(sc, x + 0.22 + index * 0.24, y + 0.18, 0.13, 0.13, fill=tone, radius=0.06, kind="deco",
                  label="dot", shape=MSO_AUTO_SHAPE_TYPE.OVAL)
    paras = []
    for index, raw in enumerate(code.splitlines()):
        tone = "light_muted" if raw.lstrip().startswith("#") else "light"
        if raw.startswith("INI_SECTION") or raw.startswith(">>>END"):
            tone = "gold"
        paras.append(Para([Run(raw if raw else " ", "mono", 13, tone)]))
    text(sc, x + 0.25, y + 0.42, w - 0.5, h - 0.56, paras, label="code", parent=parent, wrap=False, min_scale=0.5)


def right_panel(sc, section, x, y, w, h):
    if section.code:
        code_panel(sc, section.code, x, y, w, h)
    elif section.table:
        t = section.table
        table_block(sc, t["columns"], t["rows"], t["widths"], x=x, y=y, w=w, h=h, size=12, align=t.get("align"))
    elif section.visual == "bars_models":
        visual_bars(sc, [(name, count) for name, count in sc_facts["model_groups"]], x=x, y=y, w=w, h=h,
                    title="Settings")
    elif section.visual == "chain":
        data = section.visual_data
        visual_chain_vertical(sc, data["chain"], data.get("caption", ""), x, y, w, h)
    elif section.body:
        parent = shape_box(sc, x, y, w, h, fill="panel", line="edge", radius=0.12, label="prose panel")
        paras = [Para(marked(t, "sans", 13.5, "light", code_color="gold"), space_before=10) for t in section.body]
        text(sc, x + 0.3, y + 0.25, w - 0.6, h - 0.5, paras, label="prose", parent=parent, min_scale=0.62)
    else:
        return


def visual_chain_vertical(sc, items, caption, x, y, w, h):
    n = len(items)
    gap = 0.12
    cap_h = 0.5 if caption else 0.0
    bh = (h - cap_h - gap * (n - 1)) / n
    for index, name in enumerate(items):
        by = y + index * (bh + gap)
        tone = "jade" if index in (0, n - 1) else "stone"
        parent = shape_box(sc, x, by, w, bh, fill=tone, line="jade" if tone == "stone" else None, radius=0.2,
                           label="vchain pill")
        text(sc, x + 0.2, by + 0.02, w - 0.4, bh - 0.04, [Para([Run(f"{index + 1}   ", "serif-bold", 15,
                                                                    "obsidian" if tone == "jade" else "copper"),
                                                                Run(name, "sans-semibold", 15,
                                                                    "obsidian" if tone == "jade" else "ivory")])],
             anchor="middle", label="vchain name", parent=parent, min_scale=0.6)
    if caption:
        text(sc, x, y + h - cap_h + 0.08, w, cap_h - 0.08, one(caption, "sans-italic", 12, "light_muted"),
             label="vchain caption", min_scale=0.6)


def callout_slide(deck, chapter, section, notes, label):
    title, body = section.callout
    sc = deck.new("content", notes=notes)
    deck.header(sc, section.kicker, title, "", chapter.accent)
    parent = shape_box(sc, MX, CONTENT_TOP - 0.3, CONTENT_W, 3.6, fill="panel", line=chapter.accent, radius=0.16,
                       line_w=1.5, label="callout")
    shape_box(sc, MX, CONTENT_TOP - 0.3, 0.12, 3.6, fill=chapter.accent, radius=0.0, kind="deco", label="bar",
              shape=MSO_AUTO_SHAPE_TYPE.RECTANGLE)
    text(sc, MX + 0.6, CONTENT_TOP - 0.05, CONTENT_W - 1.2, 3.1,
         [Para(marked(body, "sans-light", 26, "ivory", bold_font="sans-semibold", code_color="gold"))],
         anchor="middle", label="callout text", parent=parent, min_scale=0.55)
    text(sc, MX, 5.72, CONTENT_W, 0.5, one(f"From “{plain(section.title)}”", "sans-italic", 13,
                                          "light_muted"), label="callout source")
    deck.chrome(sc, label)


sc_facts: dict = {}


# ── Tree appendix ───────────────────────────────────────────────────
TREE_PT = 8.0
TREE_COLS = 3


def tree_slides(deck: Deck, tree_text: str):
    col_gap = 0.28
    col_w = (CONTENT_W - col_gap * (TREE_COLS - 1)) / TREE_COLS
    avail = (col_w - 2 * INSET_X) * 72 * WIDTH_SAFETY * 0.99
    lines: list[str] = []
    for raw in tree_text.splitlines():
        if measure(raw, "mono", TREE_PT) <= avail:
            lines.append(raw)
            continue
        marker = max(raw.find("├── "), raw.find("└── "))
        guide = raw[:marker] + ("│   " if raw[marker] == "├" else "    ")
        head = raw
        while measure(head, "mono", TREE_PT) > avail:
            cut = len(head) - 1
            while cut > marker + 8 and measure(head[:cut], "mono", TREE_PT) > avail:
                cut -= 1
            lines.append(head[:cut])
            head = guide + "  » " + head[cut:]
        lines.append(head)
    top, bottom = 1.28, 6.84
    per_col = int(((bottom - top - 2 * INSET_Y) * 72 - 1) // (TREE_PT * LINE))
    per_slide = per_col * TREE_COLS
    total = math.ceil(len(lines) / per_slide)
    for page in range(total):
        chunk = lines[page * per_slide:(page + 1) * per_slide]
        sc = deck.new("appendix", notes="Appendix A: the complete tracked file tree, continued.")
        text(sc, MX, 0.42, 9.5, 0.36, one(f"APPENDIX A  ·  COMPLETE TRACKED FILE TREE  ·  {page + 1} / {total}",
                                          "sans-semibold", 11, "copper", spacing=2.0), label="tree kicker")
        text(sc, MX, 0.80, CONTENT_W, 0.42, one(f"{deck.facts['tracked_count']:,} tracked files at commit "
                                                f"{deck.facts['head_short']}; directories in jade, files in ivory.",
                                                "sans", 13, "light_muted"), label="tree lede")
        for col in range(TREE_COLS):
            part = chunk[col * per_col:(col + 1) * per_col]
            if not part:
                continue
            x = MX + col * (col_w + col_gap)
            paras = [Para([Run(item, "mono", TREE_PT, "jade" if item.endswith("/") else "light")]) for item in part]
            text(sc, x, top, col_w, bottom - top, paras, label="tree column", wrap=False, min_scale=1.0)
        deck.chrome(sc, "Appendix A · File tree")


def closing_slide(deck: Deck):
    f = deck.facts
    sc = deck.new("divider", notes="Tlamatini — one who knows. Created by Angela López Mendoza.")
    portrait = circular_portrait(ring="copper")
    picture(sc, portrait, 1.1, 1.35, 4.5, 4.5, label="closing portrait")
    text(sc, 6.3, 1.7, 6.3, 1.2, one("Tlamatini", "serif", 64, "ivory"), label="closing title")
    text(sc, 6.33, 2.95, 6.2, 0.55, one("one who knows", "serif-italic", 24, "jade"), label="closing tagline")
    text(sc, 6.33, 3.62, 6.3, 1.6, [Para(marked(
        "Conceived, designed and built by **Angela López Mendoza**. Every number in this deck was counted "
        "from the source tree when it was generated.", "sans-light", 17, "light", bold_font="sans-semibold",
        bold_color="ivory"))], label="closing text")
    text(sc, 6.33, 5.35, 6.0, 0.4, one("Regenerate with  refresh_project_docs.py", "mono", 13, "gold"),
         label="closing code")
    text(sc, 6.33, 5.8, 6.0, 0.36, one(f"{f['generated_date']}  ·  commit {f['head_short']}  ·  "
                                       "github.com/XAIHT/Tlamatini", "sans", 11, "light_muted"), label="closing meta")


# ── Audit ───────────────────────────────────────────────────────────
def _overlap(a: Rect, b: Rect) -> float:
    dx = min(a.x + a.w, b.x + b.w) - max(a.x, b.x)
    dy = min(a.y + a.h, b.y + b.h) - max(a.y, b.y)
    return dx * dy if dx > 0.004 and dy > 0.004 else 0.0


def audit(canvases: list[SlideCanvas]) -> list[str]:
    problems = []
    for sc in canvases:
        for rect in sc.rects:
            if rect.kind == "bg":
                continue
            if rect.x < -0.001 or rect.y < -0.001 or rect.x + rect.w > SW + 0.001 or rect.y + rect.h > SH + 0.001:
                problems.append(f"slide {sc.number}: {rect.label} leaves the slide")
            if rect.parent is not None:
                p = sc.rects[rect.parent]
                if rect.x < p.x - 0.001 or rect.y < p.y - 0.001 or rect.x + rect.w > p.x + p.w + 0.001 \
                        or rect.y + rect.h > p.y + p.h + 0.001:
                    problems.append(f"slide {sc.number}: {rect.label} escapes its {p.label}")
        texts = [r for r in sc.rects if r.kind == "text"]
        for i, a in enumerate(texts):
            for b in texts[i + 1:]:
                if _overlap(a, b):
                    problems.append(f"slide {sc.number}: text boxes overlap ({a.label} / {b.label})")
        blocks = [r for r in sc.rects if r.kind in ("image", "chart")]
        for a in texts:
            for b in blocks:
                if _overlap(a, b):
                    problems.append(f"slide {sc.number}: text overlaps {b.label} ({a.label})")
    return problems


def build_pptx(facts: dict, chapters, output) -> int:
    sc_facts.clear()
    sc_facts.update(facts)
    deck = Deck(facts)
    cover_slide(deck)
    agenda_slide(deck, chapters)
    for chapter in chapters:
        sc = divider_slide(deck, chapter)
        deck.chrome(sc, f"Chapter {chapter.numeral}")
        for section in chapter.sections:
            if section.deck != "none":
                section_slides(deck, chapter, section)
    from dossier_content import Chapter

    appendix = Chapter("tree", "A", "The complete tracked file tree",
                       f"Every one of the {facts['tracked_count']:,} files Git tracks, drawn as a tree.", [],
                       accent="gold")
    sc = divider_slide(deck, appendix)
    deck.chrome(sc, "Appendix A")
    tree_slides(deck, facts["tree_text"])
    closing_slide(deck)
    problems = audit(deck.canvases)
    if problems:
        raise LayoutError("Deck geometry audit failed:\n  " + "\n  ".join(problems[:200]))
    LEDGER.assert_clean("PPTX")
    core = deck.prs.core_properties
    core.author = "Angela López Mendoza"
    core.title = "Tlamatini — Complete Project Dossier"
    core.subject = f"Tlamatini {facts['version']}"
    core.keywords = "Tlamatini; agents; Multi-Turn; ACPX; MCP; workflow designer"
    core.comments = "Generated by refresh_project_docs.py from the source tree."
    deck.prs.save(str(output))
    return len(deck.canvases)

