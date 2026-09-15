"""Regression tests for real PPTX content preservation and bounded text.

For native PowerPoint evidence, run scripts/verify_pptxer_layout.py --render.
These tests require actual installed font files and python-pptx.
"""
import copy
from pathlib import Path
import sys
import tempfile
import unittest
import os
import subprocess
from unittest.mock import patch

AGENT_DIR = Path(__file__).parent / "agents/pptxer"
sys.path.insert(0, str(AGENT_DIR))
import pptxer_audit as audit
import pptxer_build as build
import pptxer_docmodel as model
import pptxer_fonts as fonts
import pptxer_nuance as nuance
import pptxer_theme as themes
import pptxer_render as render
from pptx import Presentation
from pptx.util import Inches, Pt


def long_text(key):
    tail = f" END_{key}"
    return ("Wide windows preserve every word with measured wrapping. " * 8)[:256 - len(tail)] + tail


class TextMeasurementRegressionTests(unittest.TestCase):
    def test_small_accent_labels_have_body_text_contrast(self):
        from pptxer_color import contrast_ratio
        theme = themes.build_theme(nuance.classify("", requested="corporate"), {})
        for role in ("kicker", "accent_text"):
            self.assertGreaterEqual(contrast_ratio(theme.color(role), theme.color("ground_worst")), 7)
        self.assertGreaterEqual(contrast_ratio(theme.color("accent_text"), theme.color("surface_high")), 7)

    @classmethod
    def setUpClass(cls):
        cls.inventory = fonts.get_inventory()
        cls.font = cls.inventory.resolve("Arial", "sans_body")

    def test_fractional_floor_is_measured_at_the_returned_size(self):
        text = long_text("floor")
        size, lines = fonts.fit_point_size(text, self.font, 52, 1, 14.2, min_pt=10.1,
                                           break_long_words=True)
        self.assertEqual(size, 10.1)
        self.assertEqual(lines, fonts.wrap_text_to_width(text, self.font, size, 52,
                                                        break_long_words=True))

    def test_start_below_floor_never_returns_empty_content(self):
        size, lines = fonts.fit_point_size("visible", self.font, 5, 5, 7, min_pt=11)
        self.assertEqual(size, 11)
        self.assertEqual(lines, ["visible"])

    def test_tracking_and_bold_are_part_of_measurement(self):
        theme = themes.build_theme(nuance.classify("", requested="corporate"), {})
        font = theme.font_for("stat_label")
        self.assertTrue(font.bold)
        self.assertEqual(font.tracking_em, theme.letter_spacing("stat_label"))
        plain = copy.copy(font)
        plain.tracking_em = 0
        self.assertGreater(fonts.measure_text("WIDE LABEL", font, 20)[0],
                           fonts.measure_text("WIDE LABEL", plain, 20)[0])

    def test_boundary_lengths_are_wrapped_without_losing_characters(self):
        for count in (255, 256, 257, 1024):
            with self.subTest(count=count):
                lines = fonts.wrap_text_to_width("W" * count, self.font, 14, 180,
                                                 break_long_words=True)
                self.assertEqual("".join(lines), "W" * count)
                self.assertTrue(all(fonts.measure_text(line, self.font, 14)[0] <= 180 for line in lines))

    def test_indentation_and_combining_marks_survive_wrapping(self):
        text = "    value = '" + "a\u0301" * 128 + "'"
        lines = fonts.wrap_text_to_width(text, self.font, 14, 180, break_long_words=True)
        self.assertEqual("".join(lines), text)
        self.assertTrue(lines[0].startswith("    "))
        self.assertTrue(all(not line.startswith("\u0301") for line in lines))


class SavedDeckAuditRegressionTests(unittest.TestCase):
    def test_unreadable_deck_is_not_reported_clean(self):
        report = audit.audit_pptx("missing-visibility-test.pptx")
        self.assertFalse(report.layout_clean)
        self.assertTrue(report.measurement_errors)
        self.assertIn("AUDIT INCOMPLETE", report.summary())

    def test_preview_is_not_called_an_independent_renderer(self):
        report = audit.AuditReport()
        report.pixel_audited = 1
        report.renderer_tier = "geometric_preview"
        self.assertIn("approximate", report.confidence)

    def test_table_cell_overflow_is_detected(self):
        prs = Presentation()
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        table = sl.shapes.add_table(1, 1, Inches(1), Inches(1), Inches(2), Inches(.4)).table
        table.cell(0, 0).text = long_text("table")
        for para in table.cell(0, 0).text_frame.paragraphs:
            for run in para.runs:
                run.font.name, run.font.size = "Arial", Pt(24)
        result = audit.audit_text_fit(prs)
        self.assertEqual(result["error"], "")
        self.assertEqual(result["frames"], 1)
        self.assertEqual(len(result["overflows"]), 1)
        self.assertIn("[1,1]", result["overflows"][0]["name"])

    def test_nonwrapping_frame_reports_horizontal_overflow(self):
        prs = Presentation()
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        shape = sl.shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(6))
        shape.text_frame.word_wrap = False
        shape.text = "W" * 256
        shape.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
        result = audit.audit_text_fit(prs)
        self.assertEqual(len(result["overflows"]), 1)
        self.assertIn("width excess", result["overflows"][0]["reason"])

    def test_soft_breaks_and_paragraph_spacing_are_counted(self):
        prs = Presentation()
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        shape = sl.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        shape.text = "one\vtwo\vthree"
        p = shape.text_frame.paragraphs[0]
        p.font.size = Pt(16)
        p.line_spacing, p.space_after = Pt(20), Pt(30)
        self.assertTrue(audit.audit_text_fit(prs)["overflows"])

    def test_text_inside_text_is_still_an_overlap(self):
        prs = Presentation()
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        for width in (4, 2):
            shape = sl.shapes.add_textbox(Inches(1), Inches(1), Inches(width), Inches(1))
            shape.text = "Overlaid content"
        self.assertEqual(len(audit.audit_shape_tree(prs)["overlaps"]), 1)

    def test_content_outside_safe_area_counts_as_a_defect(self):
        report = audit.AuditReport()
        report.safe_escapes = [{"decorative": False}]
        self.assertFalse(report.layout_clean)

    def test_grouped_text_is_audited(self):
        prs = Presentation()
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        group = sl.shapes.add_group_shape()
        shape = group.shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(.2))
        shape.text = long_text("group")
        self.assertTrue(audit.audit_text_fit(prs)["overflows"])


class ContentPaginationRegressionTests(unittest.TestCase):
    def test_portrait_cards_keep_short_values_on_one_line(self):
        theme = themes.build_theme(nuance.classify("", requested="corporate"), {})
        deck = model.Deck([model.Slide(kind="stats", title="Measured values", stats=[
            {"value": "142%", "label": long_text(f"portrait{i}")} for i in range(6)])])
        with tempfile.TemporaryDirectory() as folder:
            path = str(Path(folder) / "portrait.pptx")
            result = build.build_deck(deck, theme, path, {"generate_art": False, "slide_size": "vertical"})
            self.assertTrue(result.ok, result.error)
            values = [shape.text for slide in Presentation(path).slides for shape in slide.shapes
                      if shape.name == "pptxer:statvalue"]
            self.assertEqual(values, ["142%"] * 6)
            self.assertEqual(result.placement_diagnostics, [])

    def test_quote_reserves_space_for_attribution_and_footer(self):
        theme = themes.build_theme(nuance.classify("", requested="cyberpunk"), {})
        deck = model.Deck([model.Slide(kind="quote", quote=long_text("quote"),
                                     attribution=long_text("attribution"))])
        with tempfile.TemporaryDirectory() as folder:
            path = str(Path(folder) / "quote.pptx")
            result = build.build_deck(deck, theme, path, {"generate_art": False})
            self.assertTrue(result.ok, result.error)
            self.assertEqual(result.placement_diagnostics, [])
            self.assertEqual(audit.audit_shape_tree(path, tolerance_pt=0)["overlaps"], [])

    def test_pie_and_doughnut_keep_distinct_slices_and_complete_label_keys(self):
        theme = themes.build_theme(nuance.classify("", requested="corporate"), {})
        labels = [long_text(f"category{i}") for i in range(3)]
        deck = model.Deck([model.Slide(kind="chart", title=kind, chart={
            "kind": kind, "categories": labels,
            "series": [{"name": "Measured values", "values": [12, 25, 18]}],
        }) for kind in ("pie", "doughnut")])
        with tempfile.TemporaryDirectory() as folder:
            path = str(Path(folder) / "charts.pptx")
            result = build.build_deck(deck, theme, path, {"generate_art": False})
            self.assertTrue(result.ok, result.error)
            prs = Presentation(path)
            charts = [shape.chart for slide in prs.slides for shape in slide.shapes if shape.has_chart]
            self.assertEqual(len(charts), 2)
            for chart in charts:
                self.assertTrue(chart.has_legend)
                self.assertEqual(list(chart.series[0].values), [12, 25, 18])
                colors = {str(point.format.fill.fore_color.rgb) for point in chart.series[0].points}
                self.assertEqual(len(colors), 3)
            text = "".join(frame.text for slide in prs.slides for _, frame, _, _ in audit._text_frames(slide.shapes))
            normalized = "".join(text.split()).casefold()
            for label in labels:
                self.assertIn("".join(label.split()).casefold(), normalized)

    def test_parsed_continuations_preserve_speaker_notes_and_metadata(self):
        slide = model.Slide(title="Evidence", bullets=["item"] * 8,
                            notes="Speaker evidence", meta={"source": "fixture"})
        deck = model._refine(model.Deck([slide]), max_bullets=3)
        self.assertEqual(sum(len(s.bullets) for s in deck.slides), 8)
        self.assertTrue(all(s.notes == "Speaker evidence" and s.meta["source"] == "fixture" for s in deck.slides))

    def test_long_cards_columns_and_table_are_complete_and_bounded(self):
        sources = [long_text(f"item{i}") for i in range(12)]
        deck = model.Deck([
            model.Slide(kind="stats", title="Cards", stats=[{"value": str(i + 1), "label": text}
                       for i, text in enumerate(sources[:8])]),
            model.Slide(kind="two_column", title="Columns", columns=[
                {"title": "First", "bullets": sources[8:10]},
                {"title": "Second", "bullets": sources[10:]}]),
            model.Slide(kind="table", title="Evidence", table={"headers": ["Index", "Text"],
                       "rows": [[str(i), text] for i, text in enumerate(sources)]}),
            model.Slide(kind="code", title="Code", code="def test():\n    return '" + "W" * 256 + "'"),
        ])
        before = copy.deepcopy(deck)
        theme = themes.build_theme(nuance.classify("", requested="corporate"), {})
        with tempfile.TemporaryDirectory() as folder:
            path = str(Path(folder) / "stress.pptx")
            result = build.build_deck(deck, theme, path, {"generate_art": False, "slide_size": "4:3"})
            self.assertTrue(result.ok, result.error)
            self.assertGreater(result.slide_count, len(deck.slides))
            self.assertEqual(result.placement_diagnostics, [])
            self.assertEqual(audit.audit_text_fit(path)["overflows"], [])
            geometry = audit.audit_shape_tree(path)
            self.assertEqual(geometry["overlaps"] + geometry["off_slide"] + geometry["safe_escapes"], [])
            frames = [frame.text for slide in Presentation(path).slides
                      for _, frame, _, _ in audit._text_frames(slide.shapes)]
            all_text = "".join("".join(frames).split()).casefold()
            for text in sources:
                self.assertIn("".join(text.split()).casefold(), all_text)
            self.assertIn("w" * 256, all_text)
            self.assertTrue(any("    return" in text for text in frames))
        self.assertEqual([s.as_dict() for s in deck.slides], [s.as_dict() for s in before.slides])

    def test_impossible_page_does_not_replace_existing_file(self):
        theme = themes.build_theme(nuance.classify("", requested="corporate"), {})
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "keep.pptx"
            path.write_bytes(b"original")
            result = build.build_deck(model.Deck([model.Slide(title="X", body="Y")]), theme, str(path),
                                      {"margin_pt": 400, "generate_art": False})
            self.assertFalse(result.ok)
            self.assertEqual(path.read_bytes(), b"original")


class NativePowerPointRegressionTests(unittest.TestCase):
    def test_unknown_process_ownership_never_terminates_powerpoint(self):
        with patch.object(render, "_powerpoint_pids", return_value={12345}) as snapshot:
            self.assertEqual(render._reap_orphan_powerpoint(None, grace=0, owned_pids={12345}), [])
            self.assertEqual(render._reap_orphan_powerpoint(set(), grace=0), [])
            snapshot.assert_not_called()

    def test_worker_timeout_returns_failure_instead_of_a_clean_render(self):
        with tempfile.TemporaryDirectory() as folder:
            with patch.object(render, "_powerpoint_available", return_value=(True, "test")), \
                    patch.object(render, "_temp_dir", return_value=folder), \
                    patch.object(render.subprocess, "run", side_effect=subprocess.TimeoutExpired("office", 1)):
                result = render.render_with_powerpoint("missing.pptx", folder, timeout=1)
            self.assertFalse(result.ok)
            self.assertFalse(result.is_ground_truth)
            self.assertIn("timed out", result.error)

    def test_256_character_text_stays_visible_in_powerpoint(self):
        if not render.probe_renderers()[render.TIER_POWERPOINT]["available"]:
            self.skipTest("native PowerPoint is not installed")
        theme = themes.build_theme(nuance.classify("", requested="cyberpunk"), {})
        deck = model.Deck([
            model.Slide(kind="title_slide", title=long_text("cover"), kicker=long_text("kicker"), subtitle=long_text("subtitle")),
            model.Slide(kind="image_right", title="Text and image", body=long_text("side")),
            model.Slide(kind="table", title="Table", table={"headers": ["Text"], "rows": [[long_text("cell")]]}),
        ])
        with tempfile.TemporaryDirectory() as folder:
            path = str(Path(folder) / "native.pptx")
            result = build.build_deck(deck, theme, path, {"generate_art": False, "slide_size": "vertical"})
            self.assertTrue(result.ok, result.error)
            previous = os.environ.get("TLAMATINI_TEMP")
            os.environ["TLAMATINI_TEMP"] = folder
            try:
                rendered = render.render_slides(path, str(Path(folder) / "slides"), prefer="powerpoint")
            finally:
                if previous is None:
                    os.environ.pop("TLAMATINI_TEMP", None)
                else:
                    os.environ["TLAMATINI_TEMP"] = previous
            self.assertTrue(rendered["ok"], rendered)
            report = audit.audit_pptx(path, rendered)
            self.assertTrue(report.ground_truth, report.as_dict())
            self.assertEqual(report.defect_count, 0, report.as_dict())
            self.assertGreater(len(rendered["text_bounds"]["frames"]), 0)


if __name__ == "__main__":
    unittest.main()
