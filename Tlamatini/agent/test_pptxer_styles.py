"""Style selection, user overrides, and long-content layout regressions."""
import copy
from pathlib import Path
import sys
import tempfile
import unittest

sys.path.insert(0, str(Path(__file__).parent / "agents/pptxer"))
import pptxer_audit as audit
import pptxer_build as build
import pptxer_color as color
import pptxer_docmodel as model
import pptxer_fonts as fonts
import pptxer_nuance as nuance
import pptxer_theme as themes
from pptx import Presentation


class PresentationStyleTests(unittest.TestCase):
    def test_long_courier_line_does_not_accumulate_advance_rounding(self):
        inventory = fonts.get_inventory()
        if not inventory.has("Courier New"):
            self.skipTest("Courier New is not installed")
        font = inventory.resolve("Courier New", "mono")
        # Courier New's fixed advance is 0.6 em. The native PowerPoint
        # regression measured 837.38pt; coarse sampling returned 826.5pt.
        self.assertGreaterEqual(fonts.measure_text("W" * 87, font, 16)[0], 87 * 16 * .6)
        lines = fonts.wrap_text_to_width("W" * 256, font, 16, 829.6, break_long_words=True)
        self.assertEqual("".join(lines), "W" * 256)
        self.assertTrue(all(len(line) * 16 * .6 <= 829.6 for line in lines))

    def test_compressed_font_is_not_measured_as_regular_bodoni(self):
        self.assertEqual(fonts._fold_width_into_family("Bodoni MT", "Poster Compressed"),
                         ("Bodoni MT Poster Compressed", "regular"))
        inventory = fonts.get_inventory()
        if inventory.has("Bodoni MT"):
            from PIL import ImageFont
            resolved = inventory.resolve("Bodoni MT", "display_elegant")
            family, style = ImageFont.truetype(resolved.path, 16).getname()
            self.assertEqual(family, "Bodoni MT")
            self.assertNotIn("compressed", style.lower())

    def test_all_twelve_styles_are_selectable_by_name_and_label(self):
        self.assertEqual(len(nuance.STYLE_PRESETS), 12)
        for key, spec in nuance.STYLE_PRESETS.items():
            for name in (key, spec["label"], spec["label"] + " style"):
                with self.subTest(name=name):
                    verdict = nuance.classify("quarterly revenue report", requested=name)
                    self.assertEqual(verdict.nuance, key)
                    self.assertEqual(verdict.confidence, 1.0)

    def test_style_defaults_are_distinct_and_do_not_mutate(self):
        before = copy.deepcopy(nuance.STYLE_PRESETS)
        signatures = set()
        for key in nuance.STYLE_PRESETS:
            theme = themes.build_theme(nuance.classify("", requested=key))
            signatures.add((theme.hex("ground"), theme.hex("accent"),
                            theme.fonts["display"].family, theme.shape["radius"],
                            theme.space("margin")))
        self.assertEqual(len(signatures), 12)
        self.assertEqual(before, nuance.STYLE_PRESETS)

    def test_user_seed_and_background_mode_override_style_palette(self):
        for key in nuance.STYLE_PRESETS:
            with self.subTest(style=key):
                verdict = nuance.classify("", requested=key)
                config = {"predominant_color": "#A83CC0", "background_color": "#FFFFFF",
                          "font_pairing": "technical", "density": "high"}
                theme = themes.build_theme(verdict, config)
                self.assertEqual(theme.hex("ground"), "FFFFFF")
                self.assertEqual(theme.meta["seed"].lower(), "#a83cc0")
                self.assertEqual(theme.meta["pairing"], "technical")
                self.assertEqual(theme.density, "high")
                self.assertEqual(config["background_color"], "#FFFFFF")
                flipped = "light" if nuance.STYLE_PRESETS[key]["background_mode"] == "dark" else "dark"
                self.assertEqual(themes.build_theme(verdict, {"background_mode": flipped}).background_mode, flipped)

    def test_small_text_contrast_in_all_styles(self):
        for key in nuance.STYLE_PRESETS:
            theme = themes.build_theme(nuance.classify("", requested=key))
            for role in ("text", "caption", "footer", "kicker", "accent_text"):
                with self.subTest(style=key, role=role):
                    self.assertGreaterEqual(color.contrast_ratio(theme.color(role), theme.color("ground_worst")), 7)
            for role, background in (("on_accent", "accent"), ("table_head_text", "table_head_bg")):
                with self.subTest(style=key, role=role):
                    self.assertGreaterEqual(color.contrast_ratio(theme.color(role), theme.color(background)), 7)

    def test_long_cards_survive_every_style_and_format(self):
        text = ("Complete labels need careful measurement and enough room to breathe. " * 5)[:244] + " END_VISIBLE"
        self.assertEqual(len(text), 256)
        with tempfile.TemporaryDirectory() as folder:
            for key in nuance.STYLE_PRESETS:
                theme = themes.build_theme(nuance.classify("", requested=key))
                for size in ("16:9", "4:3", "vertical"):
                    with self.subTest(style=key, size=size):
                        deck = model.Deck([model.Slide(kind="stats", title="Complete labels",
                            stats=[{"value": f"{i+1}42%", "label": text} for i in range(6)])])
                        path = str(Path(folder) / "cards.pptx")
                        result = build.build_deck(deck, theme, path, {"slide_size": size, "generate_art": False})
                        self.assertTrue(result.ok, result.error)
                        self.assertFalse(result.placement_diagnostics)
                        report = audit.audit_pptx(path)
                        self.assertTrue(report.layout_clean, report.as_dict())
                        frames = [shape.text for slide in Presentation(path).slides for shape in slide.shapes if shape.has_text_frame]
                        saved = "".join("".join(frames).split()).casefold()
                        self.assertEqual(saved.count("".join(text.split()).casefold()), 6)
                        for i in range(6):
                            self.assertIn(f"{i+1}42%", frames)


if __name__ == "__main__":
    unittest.main()
